"""
Mobile controller split (09-Oct-2026) for Homeocentrum_All_New_And_Updated+APIs.xlsx.

- Rewrites the nine moved endpoints to their /api/MobilePatient and /api/MobileDoctor paths in every sheet.
- Marks the moved rows (apps, controller file, old URL still answering as a deprecated alias).
- Rebuilds the API_Audience_Map sheet from Hosting/ApiAudience.cs: mobile-only endpoints, then the shared (Common) list.
- Adds the change note to README and Mobile_Start_Here.

Safe to run again. Run it after _rebuild_api_catalog_xlsx.py / _add_developer_columns.py / _mobile_integration_sheets.py.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
XLSX = HERE / "Homeocentrum_All_New_And_Updated+APIs.xlsx"
XLSX_FALLBACK = HERE / "Homeocentrum_All_New_And_Updated+APIs_UPDATED.xlsx"
CATALOG = HERE.parent / "Homeocentrum.Niga.API" / "Hosting" / "ApiAudience.cs"
NEW_HOST = "https://devapi2.homeocentrum.com"
STAMP = "Moved 09-Oct-2026"

PATIENT_CTRL = "Mobile/MobilePatientController.cs"
DOCTOR_CTRL = "Mobile/MobileDoctorController.cs"

# (method, old path, new path, controller, auth, task / screen)
MOVES = [
    ("GET", "/api/Welcome/Patient", "/api/MobilePatient/Welcome", PATIENT_CTRL, "No token", "PAT-02.02 welcome slides"),
    ("GET", "/api/PatientPortal/Home", "/api/MobilePatient/Home", PATIENT_CTRL, "Patient token", "PAT-08.02 home dashboard"),
    ("GET", "/api/Patient/Visits/{patientAppId}", "/api/MobilePatient/Visits/{patientAppId}", PATIENT_CTRL, "Patient token", "PAT-20.02 appointment detail"),
    ("GET", "/api/Waitlist/Offers", "/api/MobilePatient/Waitlist/Offers", PATIENT_CTRL, "No token", "PAT-23.02 waitlist offer poll"),
    ("POST", "/api/Account/LoginWithOtp", "/api/MobileDoctor/LoginWithOtp", DOCTOR_CTRL, "No token", "DMO-01.02 doctor phone + OTP login"),
    ("GET", "/api/registration/countries", "/api/MobileDoctor/Registration/Countries", DOCTOR_CTRL, "No token", "Doctor registration: countries"),
    ("GET", "/api/registration/states", "/api/MobileDoctor/Registration/States", DOCTOR_CTRL, "No token", "Doctor registration: states (query countryId)"),
    ("GET", "/api/registration/districts", "/api/MobileDoctor/Registration/Districts", DOCTOR_CTRL, "No token", "Doctor registration: districts (query stateId)"),
    ("GET", "/api/registration/cities", "/api/MobileDoctor/Registration/Cities", DOCTOR_CTRL, "No token", "Doctor registration: cities (query districtId)"),
]

# Old path segments (after "api/") as regex, and the replacement segment. Lookbehinds keep a second run from
# touching paths that already start with MobilePatient/ or MobileDoctor/.
_GUARD = r"(?<![\w])(?<!MobileDoctor/)(?<!MobilePatient/)"
SWAPS = [
    (re.compile(_GUARD + r"Welcome/Patient(?![\w])", re.I), "MobilePatient/Welcome"),
    (re.compile(_GUARD + r"PatientPortal/Home(?![\w])", re.I), "MobilePatient/Home"),
    (re.compile(_GUARD + r"Patient/Visits/(?=\{|\d|:)", re.I), "MobilePatient/Visits/"),
    (re.compile(_GUARD + r"Waitlist/Offers(?![\w])", re.I), "MobilePatient/Waitlist/Offers"),
    (re.compile(_GUARD + r"Account/LoginWithOtp(?![\w])", re.I), "MobileDoctor/LoginWithOtp"),
    (re.compile(_GUARD + r"registration/countries(?![\w])", re.I), "MobileDoctor/Registration/Countries"),
    (re.compile(_GUARD + r"registration/states(?![\w])", re.I), "MobileDoctor/Registration/States"),
    (re.compile(_GUARD + r"registration/districts(?![\w])", re.I), "MobileDoctor/Registration/Districts"),
    (re.compile(_GUARD + r"registration/cities(?![\w])", re.I), "MobileDoctor/Registration/Cities"),
]

MOVED_NOTE_TEXT = re.compile(re.escape(STAMP))


def swap(text: str) -> str:
    for rx, new in SWAPS:
        text = rx.sub(lambda m: new.lower() if m.group(0).islower() else new, text)
    return text


def norm(path: str | None) -> str:
    p = (path or "").strip().split("?")[0]
    p = re.sub(r"^https?://[^/]+", "", p)
    p = re.sub(r"\{(\w+)[^}]*\}", r"{\1}", p)
    if not p.startswith("/"):
        p = "/" + p
    return p.rstrip("/").lower()


def header_map(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def moved_row(method: str | None, endpoint: str | None):
    ep = norm(re.sub(r"^(GET|POST|PUT|PATCH|DELETE)\s+", "", (endpoint or "").strip(), flags=re.I))
    m = (method or "").strip().upper()
    for move in MOVES:
        if norm(move[2]) == ep and (not m or m == move[0]):
            return move
    return None


def note_for(move) -> str:
    who = "Patient Mobile App" if move[3] == PATIENT_CTRL else "Doctor Mobile App"
    return (f"{STAMP}: this endpoint is only for the {who} and now lives in {move[3]}. "
            f"Call {move[0]} {NEW_HOST}{move[2]}. The old URL {move[0]} {move[1]} still answers with the same response "
            f"but is deprecated; switch to the new URL.")


def update_api_sheet(ws) -> int:
    """Sheets with the All_APIs layout (API Endpoint in column B)."""
    h = header_map(ws)
    ep_col, meth_col = h.get("API Endpoint (as in MD)"), h.get("Method Type")
    if not ep_col:
        return 0
    marked = 0
    for r in range(2, ws.max_row + 1):
        move = moved_row(ws.cell(r, meth_col).value if meth_col else None, ws.cell(r, ep_col).value)
        if not move:
            continue
        patient = move[3] == PATIENT_CTRL
        for name, value in (("Patient app", "Yes" if patient else "No"),
                            ("Doctor mobile app", "No" if patient else "Yes"),
                            ("Clinic web", "No"),
                            ("Common for all", "No"),
                            ("Who should call it", "Patient app" if patient else "Doctor mobile app")):
            if name in h:
                ws.cell(r, h[name]).value = value
        if "Mobile developer details" in h:
            cell = ws.cell(r, h["Mobile developer details"])
            text = str(cell.value or "")
            if not MOVED_NOTE_TEXT.search(text):
                cell.value = note_for(move) + ("\n\n" + text if text else "")
        marked += 1
    return marked


def update_contract_sheet(ws) -> int:
    h = header_map(ws)
    if "Endpoint" not in h:
        return 0
    marked = 0
    for r in range(2, ws.max_row + 1):
        move = moved_row(ws.cell(r, h["HTTP Method"]).value if "HTTP Method" in h else None, ws.cell(r, h["Endpoint"]).value)
        if not move:
            continue
        patient = move[3] == PATIENT_CTRL
        updates = {
            "Controller file": move[3],
            "Module": "Mobile Patient App" if patient else "Mobile Doctor App",
            "Patient app": "Yes" if patient else "No",
            "Doctor mobile app": "No" if patient else "Yes",
            "Clinic web": "No",
            "Deprecated": f"No. The old URL {move[0]} {move[1]} still answers and is deprecated.",
        }
        for name, value in updates.items():
            if name in h:
                ws.cell(r, h[name]).value = value
        marked += 1
    return marked


def read_catalog() -> list[tuple[str, str, str, str]]:
    src = CATALOG.read_text(encoding="utf-8")
    rows = re.findall(r'\("(GET|POST|PUT|PATCH|DELETE)",\s*"([^"]+)",\s*([WPD |]+),\s*"([^"]*)"\)', src)
    out = []
    for method, path, flags, use in rows:
        parts = flags.replace(" ", "").split("|")
        names = [n for f, n in (("W", "Web"), ("P", "Patient Mobile App"), ("D", "Doctor Mobile App")) if f in parts]
        out.append((method, "/" + path, ", ".join(names), use))
    return out


def lookups(wb):
    """Controller file and auth for an endpoint, read from the existing sheets."""
    ctrl, auth = {}, {}
    if "Mobile_API_Contract" in wb.sheetnames:
        ws = wb["Mobile_API_Contract"]
        h = header_map(ws)
        for r in range(2, ws.max_row + 1):
            key = ((ws.cell(r, h["HTTP Method"]).value or "").upper(), norm(ws.cell(r, h["Endpoint"]).value))
            if "Controller file" in h and ws.cell(r, h["Controller file"]).value:
                ctrl.setdefault(key, ws.cell(r, h["Controller file"]).value)
            if "Authentication required" in h and ws.cell(r, h["Authentication required"]).value:
                auth.setdefault(key, ws.cell(r, h["Authentication required"]).value)
    return ctrl, auth


def build_map_sheet(wb) -> None:
    name = "API_Audience_Map"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, index=1)
    ctrl, auth = lookups(wb)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    section_fill = PatternFill("solid", fgColor="FFF2CC")

    intro = [
        "Which endpoints are for which client (09-Oct-2026).",
        "Mobile Patient App only: /api/MobilePatient/*.  Mobile Doctor App only: /api/MobileDoctor/*.",
        "Common: shared by the web and a mobile app, or by both apps. These keep their existing paths and controllers.",
        "Web: every other endpoint. The mobile apps do not call them.",
        "Old URLs of the nine moved endpoints still answer with the same response. They are deprecated; switch to the new URL.",
        f"Swagger (sign in at {NEW_HOST}/swagger-login): pick a definition at the top right. "
        "'Patient Mobile App' and 'Doctor Mobile App' list everything that app calls; 'Common' and 'Web' list those groups; "
        "'All APIs' (v1) is the full list including the deprecated old URLs.",
        f"JSON: {NEW_HOST}/swagger/mobile-patient/swagger.json, /swagger/mobile-doctor/swagger.json, /swagger/common/swagger.json, /swagger/web/swagger.json, /swagger/v1/swagger.json",
    ]
    for line in intro:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])

    headers = ["Audience", "Method", "Endpoint (use this)", "Full URL", "Old URL (still answers, deprecated)",
               "Controller file", "Token", "Used by", "Task / screen", "Swagger definition"]
    ws.append(headers)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = bold
        c.fill = head_fill

    def section(title: str) -> None:
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        for col in range(1, len(headers) + 1):
            ws.cell(ws.max_row, col).fill = section_fill

    section("Mobile Patient App only")
    for m in MOVES:
        if m[3] == PATIENT_CTRL:
            ws.append(["Mobile Patient", m[0], m[2], NEW_HOST + m[2], f"{m[0]} {m[1]}", m[3], m[4],
                       "Patient Mobile App", m[5], "Patient Mobile App"])
    section("Mobile Doctor App only")
    for m in MOVES:
        if m[3] == DOCTOR_CTRL:
            ws.append(["Mobile Doctor", m[0], m[2], NEW_HOST + m[2], f"{m[0]} {m[1]}", m[3], m[4],
                       "Doctor Mobile App", m[5], "Doctor Mobile App"])
    section("Common (shared). Path unchanged.")
    for method, path, used_by, use in read_catalog():
        key = (method, norm(path))
        docs = ["Common"]
        if "Patient Mobile App" in used_by:
            docs.append("Patient Mobile App")
        if "Doctor Mobile App" in used_by:
            docs.append("Doctor Mobile App")
        ws.append(["Common", method, path, NEW_HOST + path, "", ctrl.get(key, ""), auth.get(key, ""),
                   used_by, use, ", ".join(docs)])
    section("Web: every endpoint not listed above. See the Clinic_Web sheet and the 'Web' Swagger definition.")

    widths = [16, 8, 46, 70, 44, 40, 14, 40, 52, 36]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(header_row, i).column_letter].width = w
    for row in ws.iter_rows(min_row=1, max_row=header_row - 2):
        for c in row:
            c.alignment = Alignment(wrap_text=False)
    ws.freeze_panes = ws.cell(header_row + 1, 1)


README_LINES = [
    "Changes 09-Oct-2026: mobile controllers",
    "New: MobilePatientController (/api/MobilePatient/*) and MobileDoctorController (/api/MobileDoctor/*) hold the endpoints only one mobile app calls. Shared endpoints keep their paths. Full list: API_Audience_Map.",
    "Moved, Patient app: GET /api/Welcome/Patient -> /api/MobilePatient/Welcome; GET /api/PatientPortal/Home -> /api/MobilePatient/Home; GET /api/Patient/Visits/{patientAppId} -> /api/MobilePatient/Visits/{patientAppId}; GET /api/Waitlist/Offers -> /api/MobilePatient/Waitlist/Offers.",
    "Moved, Doctor app: POST /api/Account/LoginWithOtp -> /api/MobileDoctor/LoginWithOtp; GET /api/registration/countries|states|districts|cities -> /api/MobileDoctor/Registration/Countries|States|Districts|Cities. GET /api/registration/qualifications is shared with the web and did not move.",
    "The old URLs still answer with the same response and are marked deprecated in Swagger. Request bodies, responses, auth and status codes did not change.",
    "Swagger has five definitions: All APIs (v1, unchanged URL), Patient Mobile App, Doctor Mobile App, Common, Web.",
]

START_HERE_ROW = [
    "Mobile-only endpoints",
    "/api/MobilePatient/* is only for the Patient Mobile App. /api/MobileDoctor/* is only for the Doctor Mobile App. Every other endpoint the apps call is shared and keeps its path. API_Audience_Map lists both groups. Swagger: choose 'Patient Mobile App' or 'Doctor Mobile App' at the top right.",
    "Use the new /api/MobilePatient and /api/MobileDoctor URLs. The old URLs of the moved endpoints still work but are deprecated.",
]


def update_readme(wb) -> None:
    ws = wb["README"]
    existing = [str(c.value) for c in ws["A"] if c.value]
    if README_LINES[0] in existing:
        start = next(c.row for c in ws["A"] if c.value == README_LINES[0])
        for i, line in enumerate(README_LINES):
            ws.cell(start + i, 1).value = line
    else:
        row = ws.max_row + 2
        for i, line in enumerate(README_LINES):
            ws.cell(row + i, 1).value = line
        ws.cell(row, 1).font = Font(bold=True)
    for c in ws["A"]:
        if isinstance(c.value, str) and c.value.startswith("Sheets:") and "API_Audience_Map" not in c.value:
            c.value = c.value.replace("Sheets: All_APIs,", "Sheets: All_APIs, API_Audience_Map,")


def update_start_here(wb) -> None:
    if "Mobile_Start_Here" not in wb.sheetnames:
        return
    ws = wb["Mobile_Start_Here"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == START_HERE_ROW[0]:
            for i, v in enumerate(START_HERE_ROW, start=1):
                ws.cell(r, i).value = v
            return
    ws.insert_rows(2)
    for i, v in enumerate(START_HERE_ROW, start=1):
        ws.cell(2, i).value = v
        ws.cell(2, i).alignment = Alignment(wrap_text=True, vertical="top")


def names_old_url(text: str) -> bool:
    """Text this script wrote on purpose with the old URLs in it."""
    return STAMP in text or text.startswith("No. The old URL") or text in README_LINES or text in START_HERE_ROW


def main() -> None:
    wb = load_workbook(XLSX)
    changed = 0
    for ws in wb.worksheets:
        if ws.title == "API_Audience_Map":
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not names_old_url(c.value):
                    new = swap(c.value)
                    if new != c.value:
                        c.value = new
                        changed += 1
    marked = {}
    for ws in wb.worksheets:
        h = header_map(ws)
        if "API Endpoint (as in MD)" in h:
            marked[ws.title] = update_api_sheet(ws)
        elif ws.title == "Mobile_API_Contract":
            marked[ws.title] = update_contract_sheet(ws)
    build_map_sheet(wb)
    update_readme(wb)
    update_start_here(wb)
    try:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}: {changed} cells re-pathed; moved rows marked: {marked}")
    except PermissionError:
        wb.save(XLSX_FALLBACK)
        print(f"Workbook is open. Saved {XLSX_FALLBACK.name} instead ({changed} cells).")


if __name__ == "__main__":
    main()
