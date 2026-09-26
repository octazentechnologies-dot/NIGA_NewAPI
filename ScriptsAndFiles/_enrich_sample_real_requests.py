# -*- coding: utf-8 -*-
"""
Refresh 'Sample Real request' with 3 copy-paste JSON/HTTP examples per API.
Uses concrete HomeoCentrum_Dev ids. Never injects unrelated /health calls.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent / "Homeocentrum_All_New_And_Updated+APIs.xlsx"
OUT_FALLBACK = Path(__file__).resolve().parent / "Homeocentrum_All_New_And_Updated+APIs_UPDATED.xlsx"

NEW_HOST = "http://127.0.0.1:5002"
OLD_HOST = "http://127.0.0.1:5001"

# Queried from HomeoCentrum_Dev
IDS = {
    "adminUserId": 10030,
    "doctorUserId": 10032,
    "patientUserId": 10033,
    "doctorId": 1010,
    "patientId": 3059,
    "patientId2": 1,
    "patientId3": 3,
    "patientAppId": 27351,
    "appointmentDate": "2026-09-25",
    "appointmentTime": "13:15:00",
    "subSectionId": 1,
    "subSectionId2": 2,
    "subSectionId3": 3,
    "metaphorId": 4,
    "aliasId": 1,
    "receptionStaffId": 1,
    "packageId": 1,
    "packageName": "9 Days",
    "loginDoctor": "Tufan_Doctor",
    "loginAdmin": "Tufan_Admin",
    "loginPatient": "Tufan_Patient",
    "loginReception": "Tufan_Reception",
    "password": "123456",
    "email": "tufanpowar001@gmail.com",
    "mobile": "7218995950",
}


def host_for(sheet_host: str | None) -> str:
    h = (sheet_host or "").lower()
    if "old" in h and "new" not in h:
        return OLD_HOST
    return NEW_HOST


def concrete_path(endpoint: str) -> str:
    path = re.sub(r"^(GET|POST|PUT|DELETE|PATCH)\s+", "", (endpoint or "").strip(), flags=re.I).strip()
    lower = path.lower()

    def repl_id(default: str) -> str:
        if "/aliases/" in lower:
            return str(IDS["aliasId"])
        if "/metaphors/" in lower:
            return str(IDS["metaphorId"])
        if "patientapp" in lower or "appointment" in lower:
            return str(IDS["patientAppId"])
        if "/patient/" in lower:
            return str(IDS["patientId"])
        if "reception" in lower:
            return str(IDS["receptionStaffId"])
        if "doctor" in lower:
            return str(IDS["doctorId"])
        if "package" in lower:
            return str(IDS["packageId"])
        return default

    path = re.sub(r"\{id(?::[^}]*)?\}", repl_id("1"), path, flags=re.I)
    path = re.sub(r"\{patientAppId(?::[^}]*)?\}", str(IDS["patientAppId"]), path, flags=re.I)
    path = re.sub(r"\{patientId(?::[^}]*)?\}", str(IDS["patientId"]), path, flags=re.I)
    path = re.sub(r"\{doctorId(?::[^}]*)?\}", str(IDS["doctorId"]), path, flags=re.I)
    path = re.sub(r"\{userId(?::[^}]*)?\}", str(IDS["doctorUserId"]), path, flags=re.I)
    path = re.sub(r"\{subSectionId(?::[^}]*)?\}", str(IDS["subSectionId"]), path, flags=re.I)
    path = re.sub(r"\{sectionId(?::[^}]*)?\}", "1", path, flags=re.I)
    path = re.sub(r"\{ticketId(?::[^}]*)?\}", "1", path, flags=re.I)
    path = re.sub(r"\{messageId(?::[^}]*)?\}", "1", path, flags=re.I)
    path = re.sub(r"\{refillId(?::[^}]*)?\}", "1", path, flags=re.I)
    path = re.sub(r"\{caseId(?::[^}]*)?\}", "1", path, flags=re.I)
    path = re.sub(r"\{packageId(?::[^}]*)?\}", str(IDS["packageId"]), path, flags=re.I)
    path = re.sub(r"\{receptionStaffId(?::[^}]*)?\}", str(IDS["receptionStaffId"]), path, flags=re.I)
    path = re.sub(r"\{sessionId(?::[^}]*)?\}", "11111111-1111-1111-1111-111111111111", path, flags=re.I)
    path = re.sub(r"\{[^}]+\}", "1", path)
    if not path.startswith("/"):
        path = "/" + path
    return path.replace("/api/api/", "/api/")


def j(obj) -> str:
    return json.dumps(obj, indent=2, ensure_ascii=False)


def http_block(n: int, method: str, url: str, body: dict | None, token: bool) -> str:
    lines = [f"Sample {n}:", f"{method} {url}"]
    if token:
        lines.append("Authorization: Bearer <PASTE_JWT_FROM_LOGIN>")
    if body is not None:
        lines.append("Content-Type: application/json")
        lines.append("")
        lines.append(j(body))
    return "\n".join(lines)


def login_bodies() -> list[dict]:
    return [
        {"userName": IDS["loginDoctor"], "password": IDS["password"]},
        {"userName": IDS["loginAdmin"], "password": IDS["password"]},
        {"userName": IDS["loginPatient"], "password": IDS["password"]},
    ]


def bodies_for(method: str, path: str) -> list[dict | None]:
    """Return up to 3 JSON bodies (None = no body). Always length 3."""
    m = method.upper()
    p = path.lower()

    if "account/login" in p or p.rstrip("/").endswith("/login"):
        return login_bodies()

    if "loginwithotp" in p or "otp" in p and "login" in p:
        return [
            {"userName": IDS["loginDoctor"], "otpChallengeId": 1, "otpCode": "123456"},
            {"userName": IDS["loginAdmin"], "otpChallengeId": 2, "otpCode": "654321"},
            {"userName": IDS["mobile"], "otpChallengeId": 3, "otpCode": "111222"},
        ]

    if "/approve" in p or "/reject" in p:
        return [{}, {}, {}]
    if "forgotpassword" in p:
        return [
            {"email": IDS["email"]},
            {"email": "tufanpowar001@gmail.com"},
            {"email": IDS["loginDoctor"]},
        ]

    if "resetpassword" in p:
        return [
            {"token": "PASTE_TOKEN_FROM_RESET_EMAIL", "newPassword": "NewPass@123"},
            {"token": "PASTE_TOKEN_FROM_RESET_EMAIL", "newPassword": "Homeo@2026"},
            {"token": "PASTE_TOKEN_FROM_RESET_EMAIL", "newPassword": "Clinic#12345"},
        ]

    if "otp/request" in p:
        return [
            {
                "action": "GrantCaregiver",
                "entityType": "Patient",
                "entityId": str(IDS["patientId"]),
                "destination": IDS["mobile"],
            },
            {
                "action": "GrantCaregiver",
                "entityType": "Patient",
                "entityId": str(IDS["patientId2"]),
                "destination": IDS["email"],
            },
            {
                "action": "Login",
                "entityType": "Mobile",
                "entityId": IDS["mobile"],
                "destination": IDS["mobile"],
            },
        ]

    if "receptionstaff/add" in p:
        return [
            {
                "doctorUserID": IDS["doctorUserId"],
                "userID": "clinic_reception_01",
                "password": IDS["password"],
                "fullName": "Clinic Reception One",
                "contactNumber": "9876543210",
                "emailId": "reception1@example.com",
            },
            {
                "doctorUserID": IDS["doctorUserId"],
                "userID": "clinic_reception_02",
                "password": IDS["password"],
                "fullName": "Clinic Reception Two",
                "contactNumber": "9123456780",
            },
            {
                "doctorUserID": IDS["doctorUserId"],
                "userID": "front_desk_03",
                "password": IDS["password"],
                "fullName": "Front Desk Three",
                "contactNumber": "9988776655",
                "emailId": "frontdesk@example.com",
            },
        ]

    if "receptionstaff/update" in p:
        return [
            {
                "receptionStaffID": IDS["receptionStaffId"],
                "fullName": "Updated Reception",
                "contactNumber": "9876543210",
                "emailId": "reception1@example.com",
            },
            {
                "receptionStaffID": IDS["receptionStaffId"],
                "fullName": "Updated Reception B",
                "contactNumber": "9123456780",
            },
            {
                "receptionStaffID": IDS["receptionStaffId"],
                "fullName": "Updated Reception C",
                "contactNumber": "9988776655",
                "emailId": None,
            },
        ]

    if "/metaphors" in p and m in ("POST", "PUT"):
        return [
            {
                "patientExpression": "as if ants crawling",
                "clinicalMeaning": "formication",
                "rubricMeaning": "MIND - SENSATIONS - crawling",
                "language": "en",
                "confidenceWeight": 0.85,
                "subSectionId": IDS["subSectionId"],
            },
            {
                "patientExpression": "like a band around head",
                "clinicalMeaning": "constricting sensation",
                "rubricMeaning": "HEAD - PAIN - band",
                "language": "en",
                "confidenceWeight": 0.9,
                "subSectionId": IDS["subSectionId2"],
            },
            {
                "patientExpression": "fear of darkness",
                "clinicalMeaning": "nyctophobia",
                "rubricMeaning": "MIND - FEAR - dark",
                "language": "hi",
                "confidenceWeight": 0.8,
                "subSectionId": IDS["subSectionId3"],
            },
        ]

    if "/aliases" in p and m in ("POST", "PUT"):
        return [
            {
                "subSectionId": IDS["subSectionId"],
                "aliasText": "restless legs at night",
                "language": "en",
                "aliasType": "patient_phrase",
                "weight": 1,
            },
            {
                "subSectionId": IDS["subSectionId2"],
                "aliasText": "fear of darkness",
                "language": "en",
                "aliasType": "patient_phrase",
                "weight": 0.8,
            },
            {
                "subSectionId": IDS["subSectionId3"],
                "aliasText": "weeping without cause",
                "language": "mr",
                "aliasType": "patient_phrase",
                "weight": 0.75,
            },
        ]

    if "availability" in p and m in ("PUT", "POST"):
        return [
            {"isOnline": True},
            {"isOnline": False},
            {"isOnline": True},
        ]

    if "patient" in p and ("create" in p or "save" in p or m == "POST") and "appointment" not in p:
        return [
            {
                "loggedInUser": IDS["doctorUserId"],
                "doctorID": IDS["doctorId"],
                "patientID": 0,
                "patientName": "Demo Patient One",
                "mobileNo": "9876500001",
                "gender": 0,
                "dateOfBirth": "1990-01-15",
                "address": "Pune",
                "countryId": 78,
                "stateId": 14,
            },
            {
                "loggedInUser": IDS["doctorUserId"],
                "doctorID": IDS["doctorId"],
                "patientID": 0,
                "patientName": "Demo Patient Two",
                "mobileNo": "9876500002",
                "gender": 1,
                "dateOfBirth": "1985-06-20",
                "address": "Mumbai",
                "countryId": 78,
                "stateId": 14,
            },
            {
                "loggedInUser": IDS["doctorUserId"],
                "doctorID": IDS["doctorId"],
                "patientID": IDS["patientId2"],
                "patientName": "pranav bandekar",
                "mobileNo": IDS["mobile"],
                "gender": 0,
                "dateOfBirth": "1992-03-10",
                "address": "Pune",
                "countryId": 78,
                "stateId": 14,
            },
        ]

    if "appointment" in p and m in ("POST", "PUT"):
        return [
            {
                "patientAppId": 0,
                "patientId": IDS["patientId"],
                "patientName": "Demo Patient",
                "doctorId": IDS["doctorId"],
                "appointmentDate": IDS["appointmentDate"],
                "appointmentTime": "10:00:00",
                "status": "WAITING",
                "deleteStatus": False,
                "userId": IDS["doctorUserId"],
                "visitType": "InClinic",
                "consultMode": "InClinic",
            },
            {
                "patientAppId": 0,
                "patientId": IDS["patientId2"],
                "patientName": "pranav bandekar",
                "doctorId": IDS["doctorId"],
                "appointmentDate": IDS["appointmentDate"],
                "appointmentTime": "11:30:00",
                "status": "WAITING",
                "deleteStatus": False,
                "userId": IDS["doctorUserId"],
                "visitType": "Tele",
                "consultMode": "Tele",
            },
            {
                "patientAppId": IDS["patientAppId"],
                "patientId": IDS["patientId"],
                "doctorId": IDS["doctorId"],
                "appointmentDate": IDS["appointmentDate"],
                "appointmentTime": IDS["appointmentTime"],
                "status": "WAITING",
                "userId": IDS["doctorUserId"],
                "visitType": "InClinic",
                "consultMode": "InClinic",
            },
        ]

    if m in ("GET", "DELETE"):
        return [None, None, None]

    # Generic writable body — still three distinct real-value payloads
    return [
        {"userId": IDS["doctorUserId"], "pageNumber": 1, "pageSize": 20},
        {"userId": IDS["adminUserId"], "pageNumber": 1, "pageSize": 50},
        {"userId": IDS["patientUserId"], "pageNumber": 2, "pageSize": 10},
    ]


def urls_for(method: str, path: str, host: str) -> list[str]:
    """Three concrete URLs for this endpoint (same resource family, different real params/ids)."""
    base = f"{host}{path}"
    m = method.upper()
    p = path.lower()

    if m == "GET":
        if "metaphor" in p and "{" not in (path):
            return [
                f"{host}/api/AudioCaseIntelligence/admin/metaphors?search=ants&language=en&pageNumber=1&pageSize=10",
                f"{host}/api/AudioCaseIntelligence/admin/metaphors?approvalStatus=Pending&pageNumber=1&pageSize=20",
                f"{host}/api/AudioCaseIntelligence/admin/metaphors/{IDS['metaphorId']}",
            ]
        if "alias" in p and "{" not in path:
            return [
                f"{host}/api/AudioCaseIntelligence/admin/aliases?search=restless&pageNumber=1&pageSize=10",
                f"{host}/api/AudioCaseIntelligence/admin/aliases?language=en&pageNumber=1&pageSize=20",
                f"{host}/api/AudioCaseIntelligence/admin/aliases/{IDS['aliasId']}",
            ]
        if "getpatientlist" in p or (p.rstrip("/").endswith("/patient") and m == "GET"):
            return [
                f"{host}/api/Patient/GetPatientList?userId={IDS['doctorUserId']}&pageNumber=1&pageSize=20",
                f"{host}/api/Patient/GetPatientList?userId={IDS['doctorUserId']}&pageNumber=1&pageSize=50",
                f"{host}/api/Patient/GetPatientList?userId={IDS['doctorUserId']}&pageNumber=2&pageSize=10",
            ]
        if "appointment" in p and "list" in p:
            return [
                f"{host}/api/PatientAppointment/GetAppointmentList?appointmentDate={IDS['appointmentDate']}&status=&userId={IDS['doctorUserId']}",
                f"{host}/api/PatientAppointment/GetAppointmentList?appointmentDate={IDS['appointmentDate']}&status=WAITING&userId={IDS['doctorUserId']}",
                f"{host}/api/PatientAppointment/GetAppointmentList?appointmentDate={IDS['appointmentDate']}&status=COMPLETED&userId={IDS['doctorUserId']}",
            ]
        if "doctordashboard" in p or "dashboard" in p and "count" in p:
            return [
                f"{base}?appointmentDate={IDS['appointmentDate']}&status=&userId={IDS['doctorUserId']}",
                f"{base}?appointmentDate={IDS['appointmentDate']}&status=WAITING&userId={IDS['doctorUserId']}",
                f"{base}?appointmentDate={IDS['appointmentDate']}&userId={IDS['doctorUserId']}",
            ]
        if "receptionstaff" in p and "list" in p:
            return [
                f"{host}/api/ReceptionStaff/GetReceptionStaffList?doctorUserID={IDS['doctorUserId']}&pageNumber=1&pageSize=50",
                f"{host}/api/ReceptionStaff/GetReceptionStaffList?doctorUserID={IDS['doctorUserId']}&pageNumber=1&pageSize=10",
                f"{host}/api/ReceptionStaff/GetReceptionStaffList?doctorUserID={IDS['doctorUserId']}&pageNumber=2&pageSize=10",
            ]
        if "searchrubricsbykeyword" in p or "keyword" in p:
            return [
                f"{host}/api/subsection/SearchRubricsByKeyword?keyword=MIND&pageNumber=1&pageSize=25",
                f"{host}/api/subsection/SearchRubricsByKeyword?keyword=HEAD&pageNumber=1&pageSize=25",
                f"{host}/api/subsection/SearchRubricsByKeyword?keyword=fear&pageNumber=1&pageSize=10",
            ]
        if path.rstrip("/") == "/health" or p.endswith("/health"):
            return [base, f"{OLD_HOST}/health", f"{NEW_HOST}/health"]
        # Generic GET: same path + 3 real query variants
        sep = "&" if "?" in base else "?"
        return [
            f"{base}{sep}pageNumber=1&pageSize=20",
            f"{base}{sep}pageNumber=1&pageSize=50",
            f"{base}{sep}pageNumber=2&pageSize=10",
        ]

    if m == "DELETE":
        if "alias" in p:
            return [
                f"{host}/api/AudioCaseIntelligence/admin/aliases/{IDS['aliasId']}",
                f"{host}/api/AudioCaseIntelligence/admin/aliases/2",
                f"{host}/api/AudioCaseIntelligence/admin/aliases",
            ]
        if "metaphor" in p:
            return [
                f"{host}/api/AudioCaseIntelligence/admin/metaphors/{IDS['metaphorId']}",
                f"{host}/api/AudioCaseIntelligence/admin/metaphors/2",
                f"{host}/api/AudioCaseIntelligence/admin/metaphors",
            ]
        # three real ids on same template
        u1 = base
        u2 = re.sub(r"/(\d+)(/?)$", r"/2\2", base)
        u3 = re.sub(r"/(\d+)(/?)$", r"/3\2", base)
        if u2 == u1:
            u2 = base.rstrip("/") + "/2"
            u3 = base.rstrip("/") + "/3"
        return [u1, u2, u3]

    # POST/PUT/PATCH — same URL, different JSON bodies
    return [base, base, base]


def three_real_requests(method: str, endpoint: str, host_label: str, token_required: str) -> str:
    host = host_for(host_label)
    path = concrete_path(endpoint)
    m = (method or "GET").upper()
    needs_token = str(token_required or "").strip().lower() in ("yes", "true", "1", "y")
    # public endpoints
    if "login" in path.lower() or "forgotpassword" in path.lower() or path.rstrip("/") == "/health":
        needs_token = False

    urls = urls_for(m, path, host)
    bodies = bodies_for(m, path)
    while len(urls) < 3:
        urls.append(urls[-1])
    while len(bodies) < 3:
        bodies.append(bodies[-1] if bodies else None)

    header = (
        "Copy-paste as-is (HomeoCentrum_Dev). "
        f"Login first if token needed: POST {NEW_HOST}/api/Account/Login "
        f'{{"userName":"{IDS["loginDoctor"]}","password":"{IDS["password"]}"}}. '
        f"Real ids: doctorUserId={IDS['doctorUserId']}, doctorId={IDS['doctorId']}, "
        f"patientId={IDS['patientId']}, patientAppId={IDS['patientAppId']}, "
        f"subSectionId={IDS['subSectionId']}, metaphorId={IDS['metaphorId']}, aliasId={IDS['aliasId']}.\n\n"
    )
    blocks = [
        http_block(1, m, urls[0], bodies[0], needs_token),
        http_block(2, m, urls[1], bodies[1], needs_token),
        http_block(3, m, urls[2], bodies[2], needs_token),
    ]
    return header + "\n\n".join(blocks)


def ensure_column(ws, after_header: str, new_header: str) -> int:
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if new_header in headers:
        return headers.index(new_header) + 1
    if after_header not in headers:
        raise RuntimeError(f"Sheet {ws.title}: missing column {after_header}")
    insert_at = headers.index(after_header) + 2
    ws.insert_cols(insert_at)
    cell = ws.cell(1, insert_at, new_header)
    cell.fill = PatternFill("solid", fgColor="1F4E79")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    return insert_at


def enrich_sheet(ws) -> int:
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if "Sample request" not in headers or "API Endpoint (as in MD)" not in headers:
        return 0
    col_real = ensure_column(ws, "Sample request", "Sample Real request")
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_ep = headers.index("API Endpoint (as in MD)") + 1
    col_method = headers.index("Method Type") + 1
    col_host = headers.index("Host") + 1 if "Host" in headers else None
    col_token = headers.index("Token required") + 1 if "Token required" in headers else None

    updated = 0
    for r in range(2, ws.max_row + 1):
        ep = ws.cell(r, col_ep).value
        if not ep:
            continue
        method = ws.cell(r, col_method).value or ""
        host = ws.cell(r, col_host).value if col_host else "New-API"
        token = ws.cell(r, col_token).value if col_token else "yes"
        text = three_real_requests(str(method), str(ep), str(host or ""), str(token or ""))
        cell = ws.cell(r, col_real, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        updated += 1
    ws.column_dimensions[ws.cell(1, col_real).column_letter].width = 60
    return updated


def main():
    wb = load_workbook(OUT)
    total = 0
    for name in wb.sheetnames:
        ws = wb[name]
        try:
            n = enrich_sheet(ws)
        except RuntimeError:
            continue
        if n:
            print(f"{name}: {n} rows")
            total += n
    try:
        wb.save(OUT)
        if OUT_FALLBACK.exists():
            try:
                OUT_FALLBACK.unlink()
            except OSError:
                pass
        print(f"Saved {OUT} ({total} cells updated)")
    except PermissionError:
        wb.save(OUT_FALLBACK)
        print(f"Original locked. Saved {OUT_FALLBACK} ({total} cells updated). Close the xlsx and re-run to replace the original.")


if __name__ == "__main__":
    main()
