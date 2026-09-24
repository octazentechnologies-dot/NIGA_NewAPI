# Adds Mobile_HowTo + Mobile_API_Reference to
# NIGA_NewAPI/ScriptsAndFiles/Homeocentrum_All_New_And_Updated+APIs.xlsx
#
# Sources (endpoints are never invented):
#   - New-API and Old-API *Controller.cs Http* attributes
#   - Old-API swagger JSON (.tmp/swagger_old.json) when present
#   - New-API swagger JSON (.tmp/swagger_new.json) when valid
#   - S1_Week1_API_DOC.txt / S2_Week2_API_DOC.txt (USE/HOST/TOKEN/SAMPLE)
#   - url_helper.js (real APIs after the Pranav marker)
#   - realbackend_helper.js (which Axios client / host)
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DOCS = Path(__file__).resolve().parent
XLSX = ROOT / "NIGA_NewAPI" / "ScriptsAndFiles" / "Homeocentrum_All_New_And_Updated+APIs.xlsx"
NEW_CTRL = ROOT / "NIGA_NewAPI" / "Homeocentrum.Niga.NewAPI" / "Controllers"
OLD_CTRL = ROOT / "NIGA_OldAPI" / "Homeocentrum.Niga.OldAPI" / "Controllers"
NEW_DTO = ROOT / "NIGA_NewAPI" / "Homeocentrum.Niga.NewAPI.Domain"
URL_HELPER = ROOT / "NIGAHomeopathy_UI" / "src" / "helpers" / "url_helper.js"
REAL_HELPER = ROOT / "NIGAHomeopathy_UI" / "src" / "helpers" / "realbackend_helper.js"
SWAGGER_OLD = ROOT / ".tmp" / "swagger_old.json"
SWAGGER_NEW = ROOT / ".tmp" / "swagger_new.json"
API_DOCS = [
    DOCS / "S1_Week1_API_DOC.txt",
    ROOT / "NIGA_NewAPI" / "ScriptsAndFiles" / "S2_Week2" / "S2_Week2_API_DOC.txt",
    DOCS / "S1_Week1_MOBILE_API_DOC.txt",
]

NEW_HOST_LOCAL = "http://127.0.0.1:5038"
OLD_HOST_LOCAL = "http://127.0.0.1:5001"
NEW_HOST_DEV = "https://devapi2.homeocentrum.com"
OLD_HOST_DEV = "https://devapi1.homeocentrum.com"

TODAY = date.today().strftime("%d-%b-%Y")

HTTP_ATTR = re.compile(
    r'\[Http(Get|Post|Put|Delete|Patch)(?:\(\s*(?:Name\s*=\s*"[^"]+"\s*,\s*)?(?:"([^"]*)")?[^)]*\))?\]',
    re.I,
)
ROUTE_ATTR = re.compile(r"\[Route\s*\(\s*\"([^\"]+)\"", re.I)
CLASS_RE = re.compile(
    r"(?P<attrs>(?:\[[^\]]+\]\s*)*)public\s+(?:abstract\s+|sealed\s+)?(?:partial\s+)?class\s+(?P<name>\w+Controller)\b",
    re.S,
)
PROP_RE = re.compile(
    r"(?P<pre>(?:\[[^\]]+\]\s*)*)public\s+(?:virtual\s+)?(?P<type>[\w<>,\?\[\]\.]+)\s+(?P<name>\w+)\s*\{\s*get",
    re.S,
)

MOBILE_PATHS = (
    "/welcome/patient",
    "/account/loginwithotp",
    "/account/requestotp",
    "/account/confirmmobile",
    "/account/forgotpassword",
    "/account/resetpassword",
    "/account/changepassword",
    "/account/logout",
    "/device",
    "/patientportal",
    "/patientauth",
    "/public/",
    "/family",
    "/caregiver",
    "/consent",
    "/otp",
    "/availability",
    "/profile/me",
    "/patientprofile",
)
WEB_ONLY_PATHS = (
    "/adminacl/",
    "/qualification/add",
    "/threedbodypart",
    "/repertorizationpage/centerofgravity",
    "/clipboard",
    "/clinicalquestions",
    "/patientlab",
    "/allopathicdrug",
    "/prescription",
    "/audiocasetaking",
    "/audiocaseintelligence",
    "/patientboardbackup",
    "/receptionstaff",
    "/exportcasetopdf",
    "/enquiry",
    "/securefile/",
    "/securedocument/",
    "/otp/audit",
    "/consent/adminaudit",
    "/knowledgegraph",
    "/aimonitoring",
    "/aiembedding",
)
CASE_TAKING_PATHS = (
    "/audiocasetaking",
    "/audiocaseintelligence",
    "/clipboard",
    "/repertorizationpage",
    "/clinicalquestions",
    "/prescription",
    "/casedetails",
    "/savecomplaints",
)

HEADERS = [
    "API ID",
    "Module",
    "Feature",
    "API Name",
    "API Status",
    "API Project",
    "Controller",
    "Endpoint",
    "HTTP Method",
    "API Version",
    "Existing/New/Updated",
    "Old API Endpoint",
    "New API Endpoint",
    "Authentication Required",
    "Authorization/Role",
    "Headers",
    "Path Parameters",
    "Query Parameters",
    "Request Body",
    "Request Fields",
    "Required/Optional",
    "Response Status",
    "Success Response",
    "Response Fields",
    "Error Responses",
    "Error Description",
    "Pagination",
    "Sorting",
    "Filtering",
    "Search Parameters",
    "Example Request",
    "Example Response",
    "Database/Entity",
    "React UI Usage",
    "Mobile Usage",
    "Local Host",
    "Dev Host",
    "Notes",
    "Source File",
    "Last Updated/Status",
]


def strip_comments(src: str) -> str:
    src = re.sub(r"/\*.*?\*/", "", src, flags=re.S)
    out = []
    for line in src.splitlines():
        if "//" in line:
            i = line.find("//")
            if "http://" in line[max(0, i - 8) : i + 8] or "https://" in line[max(0, i - 8) : i + 8]:
                out.append(line)
                continue
            line = line[:i]
        out.append(line)
    return "\n".join(out)


def controller_short(name: str) -> str:
    return name[:-10] if name.endswith("Controller") else name


def combine_route(ctrl_route: str, http_tmpl: str, ctrl_name: str, action: str) -> str:
    def sub(s: str) -> str:
        return (
            (s or "")
            .replace("[controller]", ctrl_name)
            .replace("[action]", action)
            .replace("[Controller]", ctrl_name)
        )

    ctrl_route = sub(ctrl_route or f"api/{ctrl_name}")
    http_tmpl = sub(http_tmpl or "")
    if http_tmpl.startswith("~/"):
        path = http_tmpl[1:]
    elif http_tmpl.startswith("/"):
        path = http_tmpl
    elif http_tmpl.lower().startswith("api/"):
        path = "/" + http_tmpl
    else:
        base = ctrl_route if ctrl_route.startswith("/") else "/" + ctrl_route
        path = base.rstrip("/") + ("/" + http_tmpl.lstrip("/") if http_tmpl else "")
    path = re.sub(r"/{2,}", "/", path)
    if not path.startswith("/"):
        path = "/" + path
    return path.rstrip("/") or "/"


def path_params(endpoint: str) -> str:
    parts = re.findall(r"\{([^}:]+)(?::[^}]+)?\}", endpoint)
    return ", ".join(p.strip() for p in parts) if parts else "—"


def module_of(endpoint: str, controller: str) -> str:
    p = endpoint.lower()
    mapping = [
        ("/account", "Auth / Account"),
        ("/users", "Users / Registration"),
        ("/registration", "Users / Registration"),
        ("/public", "Public directory / booking"),
        ("/family", "Family"),
        ("/caregiver", "Caregiver"),
        ("/consent", "Consent"),
        ("/otp", "OTP"),
        ("/device", "Push devices"),
        ("/patientportal", "Patient portal"),
        ("/patientauth", "Patient portal"),
        ("/welcome", "Welcome / languages"),
        ("/availability", "Availability / slots"),
        ("/patientappointment", "Appointments"),
        ("/profile", "Doctor profile"),
        ("/patientprofile", "Patient profile"),
        ("/patient", "Patients"),
        ("/receptionstaff", "Reception"),
        ("/adminacl", "Admin ACL"),
        ("/package", "Subscription / package"),
        ("/subscription", "Subscription / package"),
        ("/whatsapp", "WhatsApp"),
        ("/enquiry", "Enquiry / contact"),
        ("/securefile", "Secure files"),
        ("/securedocument", "Secure files"),
        ("/mastersapi", "Masters / dropdowns"),
        ("/dropdownlist", "Masters / dropdowns"),
        ("/pagination", "Pagination lists"),
        ("/remedy", "Repertory / remedy"),
        ("/rubric", "Repertory / rubric"),
        ("/section", "Repertory / section"),
        ("/materiamedica", "Materia medica"),
        ("/clipboard", "Clipboard / case-taking"),
        ("/audiocase", "Audio case-taking"),
        ("/repertorization", "Repertorization"),
        ("/clinical", "Clinical questions"),
        ("/prescription", "Prescription"),
        ("/patientlab", "Labs"),
        ("/allopathic", "Allopathic drugs"),
        ("/doctordashboard", "Doctor dashboard"),
        ("/news", "CMS / news"),
        ("/blog", "CMS / blog"),
        ("/qualification", "Qualifications"),
        ("/location", "Geo / location"),
        ("/state", "Geo / location"),
        ("/threed", "3D body"),
        ("/knowledgegraph", "AI / knowledge graph"),
        ("/aimonitor", "AI / monitoring"),
        ("/aiembed", "AI / embeddings"),
        ("/patientboard", "Patient board"),
        ("/export", "Exports"),
    ]
    for needle, label in mapping:
        if needle in p:
            return label
    return controller or "General"


def mobile_usage(endpoint: str, auth: str) -> str:
    p = endpoint.lower()
    if any(k in p for k in CASE_TAKING_PATHS):
        return "Do NOT use on Doctor mobile (case-taking is clinic web only). Not for Patient app."
    if any(k in p for k in WEB_ONLY_PATHS):
        return "Web / clinic SPA. Not a Patient-app or Doctor-app screen API."
    if any(k in p for k in MOBILE_PATHS) or "/public/" in p:
        if "/public/" in p or "/welcome/patient" in p or "/patientauth" in p:
            return "Patient app (and public website). Same JSON."
        if "/family" in p or "/caregiver" in p or "/consent" in p or "/device" in p:
            return "Patient app. Web SPA also uses the same URLs."
        if "/availability" in p or "/profile/me" in p:
            return "Doctor app + clinic web. Not case-taking."
        return "Mobile + Web (shared)."
    if "/account/login" in p:
        return "Clinic web login stays on Old-API. Patient/Doctor mobile OTP login is New-API /Account/LoginWithOtp."
    return "Available HTTP API. Confirm with feature owner before using on mobile."


def auth_headers(auth_required: str) -> str:
    if auth_required == "No":
        return "Content-Type: application/json (or multipart/form-data when uploading files)"
    return "Authorization: Bearer <JWT>; Content-Type: application/json (or multipart/form-data when uploading files)"


def errors_for(auth_required: str, roles: str) -> tuple[str, str]:
    codes = ["400 Bad Request"]
    desc = ["Validation / business-rule failure (message in body)."]
    if auth_required == "Yes":
        codes.append("401 Unauthorized")
        desc.append("Missing, invalid, or expired JWT.")
    if roles and roles not in ("Public / anonymous", "Authenticated (any valid JWT)", "Not confirmed"):
        codes.append("403 Forbidden")
        desc.append(f"Authenticated but role/policy denied ({roles}).")
    codes.append("500 Internal Server Error")
    desc.append("Unhandled server exception.")
    return "; ".join(codes), " ".join(desc)


def parse_controllers(folder: Path, project: str) -> list[dict]:
    rows = []
    if not folder.exists():
        return rows
    for path in sorted(folder.glob("*Controller.cs")):
        raw = strip_comments(path.read_text(encoding="utf-8", errors="replace"))
        for cm in CLASS_RE.finditer(raw):
            cls = cm.group("name")
            if cls == "BaseAPIController":
                continue
            attrs = cm.group("attrs") or ""
            ctrl_routes = ROUTE_ATTR.findall(attrs)
            class_anon = "AllowAnonymous" in attrs
            class_auth = "Authorize" in attrs and not class_anon
            class_roles = []
            if "ForbidMoneyRoles" in attrs:
                class_roles.append("Authenticated except Account and PharmacyPartner (403)")
            if "DoctorOnly" in attrs:
                class_roles.append("Doctor (Reception/Patient forbidden)")
            pol = re.findall(r'Authorize\s*\(\s*Policy\s*=\s*"([^"]+)"', attrs)
            class_roles.extend(pol)
            if "AdminPortal" in attrs:
                class_roles.append("AdminPortal")
            start = cm.end()
            nxt = CLASS_RE.search(raw, start)
            body = raw[start : nxt.start() if nxt else len(raw)]
            short = controller_short(cls)
            ctrl_route = ctrl_routes[0] if ctrl_routes else f"api/{short}"

            for hm in HTTP_ATTR.finditer(body):
                method = hm.group(1).upper()
                tmpl = hm.group(2) if hm.group(2) is not None else ""
                window_start = max(0, hm.start() - 500)
                window = body[window_start : hm.start()]
                # method name: first public IActionResult / Task after attribute cluster
                after = body[hm.end() : hm.end() + 800]
                mn = re.search(
                    r"public\s+(?:async\s+)?(?:virtual\s+)?[\w<>,\.\?]+\s+(\w+)\s*\((.*?)\)",
                    after,
                    re.S,
                )
                if not mn:
                    continue
                action = mn.group(1)
                sig = mn.group(2)
                attr_block = after[: mn.start()]
                nearby = window[-200:] + attr_block
                meth_anon = "AllowAnonymous" in attr_block
                meth_auth_attr = "Authorize" in attr_block
                if meth_anon:
                    is_anon = True
                elif meth_auth_attr or class_auth:
                    is_anon = False
                elif class_anon:
                    is_anon = True
                else:
                    # No global FallbackPolicy on these hosts → anonymous unless [Authorize]
                    is_anon = True
                meth_anon = is_anon
                meth_auth = not is_anon
                roles = list(class_roles)
                roles.extend(re.findall(r'Authorize\s*\(\s*Policy\s*=\s*"([^"]+)"', nearby))
                if "ForbidMoneyRoles" in nearby and "Authenticated except Account" not in " ".join(roles):
                    roles.append("Authenticated except Account and PharmacyPartner (403)")
                if "DoctorOnly" in nearby and "Doctor (Reception/Patient forbidden)" not in " ".join(roles):
                    roles.append("Doctor (Reception/Patient forbidden)")
                endpoint = combine_route(ctrl_route, tmpl, short, action)
                from_body = re.search(r"\[FromBody\]\s+([\w\.\?]+)\s+(\w+)", sig)
                from_form = "FromForm" in sig or "IFormFile" in sig
                from_query = re.findall(r"\[FromQuery\]\s+([\w\.\?]+)\s+(\w+)", sig)
                q_simple = re.findall(
                    r"(?:\[FromQuery\]\s+)?(?:int|long|string|bool|decimal|DateTime|Guid)\??\s+(\w+)\s*(?:=[^,\)]+)?",
                    sig,
                )
                consumes = "multipart/form-data" if from_form else ("application/json" if from_body and method != "GET" else "—")
                req_body = "—"
                if from_form:
                    req_body = "multipart/form-data (see Request Fields)"
                elif from_body:
                    req_body = from_body.group(1)
                elif method in ("POST", "PUT", "PATCH") and re.search(r"\b\w+Model\b|\b\w+Request\b|\b\w+Dto\b", sig):
                    tm = re.search(r"\b([\w\.]+(?:Model|Request|Dto))\b", sig)
                    req_body = tm.group(1) if tm else "Not confirmed"
                query = []
                for t, n in from_query:
                    query.append(f"{n}:{t}")
                for n in q_simple:
                    if n not in {x.split(":")[0] for x in query} and n not in ("cancellationToken",):
                        query.append(n)
                auth_req = "No" if meth_anon else "Yes"
                if not roles:
                    role_txt = "Public / anonymous" if meth_anon else "Authenticated (any valid JWT)"
                else:
                    role_txt = "; ".join(dict.fromkeys(roles))
                rows.append(
                    {
                        "project": project,
                        "controller": cls,
                        "method": method,
                        "endpoint": endpoint,
                        "action": action,
                        "auth": auth_req,
                        "roles": role_txt,
                        "request_body": req_body,
                        "consumes": consumes,
                        "query": ", ".join(query) if query else "—",
                        "source": str(path.relative_to(ROOT)).replace("\\", "/"),
                    }
                )
    return rows


def parse_swagger(path: Path, project: str) -> list[dict]:
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    except json.JSONDecodeError:
        return []
    if "paths" not in data:
        return []
    rows = []
    for pth, ops in data.get("paths", {}).items():
        if not isinstance(ops, dict):
            continue
        for method, op in ops.items():
            if method.lower() not in ("get", "post", "put", "delete", "patch"):
                continue
            if not isinstance(op, dict):
                continue
            params = op.get("parameters") or []
            path_ps, query_ps, header_ps = [], [], []
            body = "—"
            for pr in params:
                loc = pr.get("in")
                name = pr.get("name")
                req = "required" if pr.get("required") else "optional"
                schema = pr.get("type") or (pr.get("schema") or {}).get("type") or (pr.get("schema") or {}).get("$ref", "")
                if isinstance(schema, str) and schema.startswith("#/"):
                    schema = schema.rsplit("/", 1)[-1]
                item = f"{name}:{schema} ({req})"
                if loc == "path":
                    path_ps.append(item)
                elif loc == "query":
                    query_ps.append(item)
                elif loc == "header":
                    header_ps.append(item)
                elif loc == "body":
                    body = str(schema)
            rb = op.get("requestBody") or {}
            if rb:
                content = rb.get("content") or {}
                if "multipart/form-data" in content:
                    body = "multipart/form-data"
                elif "application/json" in content:
                    schema = (content["application/json"].get("schema") or {})
                    ref = schema.get("$ref") or (schema.get("items") or {}).get("$ref") or schema.get("type") or "object"
                    body = str(ref).rsplit("/", 1)[-1]
            tags = op.get("tags") or []
            summary = op.get("summary") or op.get("operationId") or ""
            responses = op.get("responses") or {}
            ok = responses.get("200") or responses.get("201") or {}
            ok_desc = ok.get("description") or "OK"
            err = []
            for code, meta in responses.items():
                if str(code).startswith("2"):
                    continue
                err.append(f"{code} {(meta or {}).get('description') or ''}".strip())
            sec = op.get("security")
            auth = "Yes" if sec else ("No" if sec == [] else "Not confirmed")
            rows.append(
                {
                    "project": project,
                    "controller": tags[0] if tags else "",
                    "method": method.upper(),
                    "endpoint": pth if pth.startswith("/") else "/" + pth,
                    "action": op.get("operationId") or "",
                    "auth": auth,
                    "roles": "Not confirmed" if auth == "Yes" else "Public / anonymous",
                    "request_body": body,
                    "consumes": "multipart/form-data" if body == "multipart/form-data" else "application/json",
                    "query": ", ".join(query_ps) if query_ps else "—",
                    "path_swagger": ", ".join(path_ps) if path_ps else "",
                    "success": ok_desc,
                    "swagger_errors": "; ".join(err) if err else "",
                    "summary": summary,
                    "source": f"swagger:{path.name}",
                }
            )
    return rows


def parse_api_docs() -> dict[tuple[str, str], dict]:
    """(METHOD, normalized path) -> enrichment from MD docs."""
    from _build_api_catalog_xlsx import parse_file  # same folder

    out: dict[tuple[str, str], dict] = {}
    mapping = [
        (API_DOCS[0], "S1_Week1"),
        (API_DOCS[1], "S2_Week2"),
    ]
    for path, week in mapping:
        if not path.exists():
            continue
        for r in parse_file(path, week):
            ep = r["endpoint"].split()[0]
            key = (r["method"].upper(), norm(ep))
            cur = out.get(key, {})
            cur.update(
                {
                    "use": r.get("use") or cur.get("use", ""),
                    "token": r.get("token") or cur.get("token", ""),
                    "host": r.get("host") or cur.get("host", ""),
                    "sample": r.get("sample") or cur.get("sample", ""),
                    "note": r.get("note") or cur.get("note", ""),
                    "developed": r.get("developed") or cur.get("developed", ""),
                    "channel": r.get("channel") or cur.get("channel", ""),
                    "user": r.get("user") or cur.get("user", ""),
                    "week": (cur.get("week", "") + "," + week).strip(","),
                    "source_doc": r.get("source") or cur.get("source_doc", ""),
                }
            )
            out[key] = cur
    return out


def norm(p: str) -> str:
    p = re.sub(r"\?.*$", "", (p or "").strip())
    p = p.split()[-1] if p else p
    if not p.startswith("/"):
        p = "/" + p
    p = re.sub(r"\{[^}]+\}", "{id}", p)
    return p.lower().rstrip("/") or "/"


def parse_url_helper() -> dict[str, str]:
    if not URL_HELPER.exists():
        return {}
    text = URL_HELPER.read_text(encoding="utf-8", errors="replace")
    marker = "actual api urls start from here"
    i = text.lower().find(marker)
    if i >= 0:
        text = text[i:]
    out = {}
    for m in re.finditer(r'export const (\w+)\s*=\s*"([^"]+)"', text):
        out[m.group(1)] = m.group(2)
    return out


def parse_react_usage(const_to_path: dict[str, str]) -> dict[str, list[str]]:
    """normalized path -> list of helper notes."""
    usage: dict[str, list[str]] = {}
    if not REAL_HELPER.exists():
        return usage
    text = REAL_HELPER.read_text(encoding="utf-8", errors="replace")
    # nigahomeoAPI.get(url.FOO  or  "/bar"
    for m in re.finditer(
        r"(nigahomeoAPI|nigahomeoMultipart|api)\.(get|post|put|delete|patch)\(\s*(url\.(\w+)|[\"']([^\"']+)[\"'])",
        text,
        re.I,
    ):
        client = m.group(1)
        method = m.group(2).upper()
        const = m.group(4)
        literal = m.group(5)
        path = const_to_path.get(const, literal or "")
        if not path:
            continue
        if not path.startswith("/"):
            path = "/" + path
        full = path if path.lower().startswith("/api") else "/api" + path
        host = "New-API (nigahomeoAPI / API_URL_NIGAHOMEOPATHY)" if "nigahomeo" in client.lower() else "Old-API (api / API_URL)"
        note = f"React realbackend_helper: {method} {full} via {host}"
        usage.setdefault(norm(full), []).append(note)
        usage.setdefault(norm(path), []).append(note)
    return usage


def load_dto_fields() -> dict[str, tuple[str, str]]:
    """class name -> (fields list, required/optional)."""
    found: dict[str, tuple[str, str]] = {}
    if not NEW_DTO.exists():
        return found
    for path in NEW_DTO.rglob("*.cs"):
        text = strip_comments(path.read_text(encoding="utf-8", errors="replace"))
        for cm in re.finditer(r"public class (\w+)\b", text):
            name = cm.group(1)
            start = cm.end()
            nxt = re.search(r"public class \w+\b", text[start:])
            body = text[start : start + (nxt.start() if nxt else len(text) - start)]
            fields, req, opt = [], [], []
            chunks = re.split(r"(?=public\s+)", body)
            for ch in chunks:
                pm = re.match(
                    r"public\s+(?:virtual\s+)?([\w<>,\?\[\]\.]+)\s+(\w+)\s*\{",
                    ch.strip(),
                )
                if not pm:
                    continue
                typ, fname = pm.group(1), pm.group(2)
                if fname in ("get", "set") or typ in ("class", "enum"):
                    continue
                required = "[Required]" in ch or "Required(" in ch
                item = f"{fname}:{typ}"
                fields.append(item)
                (req if required else opt).append(fname)
            if fields:
                found[name] = (
                    ", ".join(fields[:40]) + (" …" if len(fields) > 40 else ""),
                    ("Required: " + ", ".join(req) if req else "No [Required] attributes found")
                    + ("; Optional: " + ", ".join(opt) if opt else ""),
                )
    return found


def classify_status(project: str, method: str, endpoint: str, docs: dict) -> str:
    key = (method.upper(), norm(endpoint))
    d = docs.get(key, {})
    developed = (d.get("developed") or "").lower()
    if "new api" in developed:
        return "Newly Developed"
    if "updated existing" in developed:
        return "Updated"
    if "existing url" in developed:
        return "Existing"
    if project.startswith("New"):
        return "Newly Developed" if "/api/" in endpoint.lower() else "Existing"
    return "Existing"


def refine_new_status(endpoint: str, method: str, docs: dict, old_keys: set) -> str:
    key = (method.upper(), norm(endpoint))
    d = docs.get(key, {})
    developed = (d.get("developed") or "").lower()
    if "updated existing" in developed:
        return "Updated"
    if "new api" in developed:
        return "Newly Developed"
    if key in old_keys:
        return "Existing"
    # New-API often mirrors classic URLs (same route, different host). If the
    # same method+path exists on Old-API, it is not a newly invented URL.
    if key in old_keys:
        return "Existing"
    return "Newly Developed"


def merge_rows(ctrl: list[dict], swagger: list[dict], project: str) -> list[dict]:
    by = {}
    for r in swagger + ctrl:
        key = (r["method"].upper(), norm(r["endpoint"]))
        cur = by.get(key, {})
        cur.update({k: v for k, v in r.items() if v not in ("", None)})
        # prefer controller endpoint (keeps {id} names) when both exist
        if r in ctrl:
            cur["endpoint"] = r["endpoint"]
            cur["controller"] = r.get("controller") or cur.get("controller")
            cur["auth"] = r.get("auth") or cur.get("auth")
            cur["roles"] = r.get("roles") or cur.get("roles")
            cur["source"] = r.get("source") or cur.get("source")
        by[key] = cur
        cur["project"] = project
    return list(by.values())


def example_url(host: str, endpoint: str) -> str:
    return host.rstrip("/") + endpoint


def build_sheet_rows(
    new_rows: list[dict],
    old_rows: list[dict],
    docs: dict,
    react: dict[str, list[str]],
    dto_fields: dict[str, tuple[str, str]],
) -> list[list]:
    old_keys = {(r["method"].upper(), norm(r["endpoint"])) for r in old_rows}
    new_keys = {(r["method"].upper(), norm(r["endpoint"])) for r in new_rows}
    out = []

    def one(r: dict, project_label: str, seq: int, prefix: str) -> list:
        method = r["method"].upper()
        ep = r["endpoint"]
        key = (method, norm(ep))
        d = docs.get(key, {})
        if project_label.startswith("New"):
            status = refine_new_status(ep, method, docs, old_keys)
        else:
            status = classify_status(project_label, method, ep, docs)
            if key in new_keys and status == "Existing":
                # same URL lives on both — still Existing on Old-API
                status = "Existing"
        old_ep = f"{method} {ep}" if project_label.startswith("Old") else (f"{method} {ep}" if key in old_keys else "—")
        new_ep = f"{method} {ep}" if project_label.startswith("New") else (f"{method} {ep}" if key in new_keys else "—")
        if project_label.startswith("Old") and key not in new_keys:
            new_ep = "—"
        if project_label.startswith("New") and key not in old_keys:
            old_ep = "—"
        auth = r.get("auth") or ("No" if (d.get("token") or "").lower().startswith("no") else "Yes" if d.get("token") else "Not confirmed")
        if (d.get("token") or "").lower().startswith("no"):
            auth = "No"
        elif d.get("token") and "jwt" in (d.get("token") or "").lower():
            auth = "Yes"
        roles = r.get("roles") or d.get("user") or "Not confirmed"
        body_type = r.get("request_body") or "—"
        fields, reqopt = "Not confirmed", "Not confirmed"
        simple = re.sub(r"\?$", "", (body_type or "").split(".")[-1])
        if simple in dto_fields:
            fields, reqopt = dto_fields[simple]
        elif body_type in ("—", ""):
            fields, reqopt = "—", "—"
        query = r.get("query") or "—"
        pag = "Yes" if re.search(r"page", query, re.I) or "page" in (fields or "").lower() else "No"
        filt = "Yes" if re.search(r"q\b|search|filter|city|minfee|maxfee", (query + fields).lower()) else "No"
        sort = "Yes" if "sort" in (query + fields).lower() else "Not confirmed"
        search = query if filt == "Yes" else "—"
        err, errd = errors_for(auth, roles)
        if r.get("swagger_errors"):
            err = r["swagger_errors"]
        sample = d.get("sample") or ""
        host_local = NEW_HOST_LOCAL if project_label.startswith("New") else OLD_HOST_LOCAL
        host_dev = NEW_HOST_DEV if project_label.startswith("New") else OLD_HOST_DEV
        react_notes = react.get(norm(ep), []) or react.get(norm("/api" + ep[4:] if ep.lower().startswith("/api") else "/api" + ep), [])
        react_txt = "; ".join(dict.fromkeys(react_notes)) if react_notes else "Not called from realbackend_helper.js (may still be a public API)."
        notes = " | ".join(x for x in [d.get("use"), d.get("note"), r.get("summary")] if x)
        if not notes:
            notes = "Contract taken from controller/Swagger. Do not invent extra fields."
        success = r.get("success") or "200 OK — JSON body. Field list: see week API DOC when listed; otherwise inspect a live 200 (do not assume)."
        example_req = f"{method} {example_url(host_local, ep)}"
        if sample:
            example_req += "  SAMPLE: " + sample[:400]
        api_name = r.get("action") or ep.rsplit("/", 1)[-1]
        feature = d.get("use") or module_of(ep, r.get("controller") or "")
        version = "v1 (no /v2 prefix in routes)"
        return [
            f"{prefix}-{seq:04d}",
            module_of(ep, r.get("controller") or ""),
            feature[:180],
            api_name,
            status,
            project_label,
            r.get("controller") or "Not confirmed",
            ep,
            method,
            version,
            status,
            old_ep,
            new_ep,
            auth,
            roles,
            auth_headers(auth),
            path_params(ep) if not r.get("path_swagger") else r.get("path_swagger"),
            query,
            body_type,
            fields,
            reqopt,
            "200 (success) plus errors in Error Responses",
            success,
            "Not confirmed unless listed in week API DOC / Swagger schema — do not invent field names.",
            err,
            errd,
            pag,
            sort,
            filt,
            search,
            example_req,
            "Use live 200 body. Week docs SAMPLE/OUTPUT when present. Do not copy guessed JSON.",
            "See controller / EF entities in Source File. Not confirmed at table level unless in SQL scripts.",
            react_txt,
            mobile_usage(ep, auth),
            host_local,
            host_dev,
            notes[:500],
            r.get("source") or d.get("source_doc") or "",
            f"{TODAY} / {status}",
        ]

    nseq = oseq = 0
    for r in sorted(new_rows, key=lambda x: (x["endpoint"].lower(), x["method"])):
        nseq += 1
        out.append(one(r, "New API (nigahomeoAPI)", nseq, "NEW"))
    for r in sorted(old_rows, key=lambda x: (x["endpoint"].lower(), x["method"])):
        oseq += 1
        out.append(one(r, "Old API (classic / api)", oseq, "OLD"))
    return out


def write_howto(wb, n_new: int, n_old: int, n_rows: int):
    name = "Mobile_HowTo"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, 1)
    ws["A1"] = "HomeoCentrum — Mobile developer API reference (how to use this workbook)"
    ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
    lines = [
        "",
        f"Generated: {TODAY}. Sheet Mobile_API_Reference has {n_rows} callable HTTP endpoints ({n_new} New-API + {n_old} Old-API).",
        "",
        "WHICH HOST TO CALL",
        f"  New API (nigahomeoAPI): local {NEW_HOST_LOCAL}   dev {NEW_HOST_DEV}",
        f"  Old API (classic api):  local {OLD_HOST_LOCAL}   dev {OLD_HOST_DEV}",
        "  Clinic SPA config.js: API_URL_NIGAHOMEOPATHY = New-API /api ; API_URL = Old-API /api",
        "  Paths in the sheet already include /api/…  Do not prefix /api again.",
        "",
        "DUAL-API RULE (do not switch hosts silently)",
        "  Login (clinic username/password), Rx-write, Razorpay → Old API.",
        "  New HTTP from S1/S2 (public directory, family, caregiver, OTP, consent, patient portal,",
        "  availability, secure password, doctor registration with documents, …) → New API.",
        "  If a URL exists on BOTH projects, keep using the host the live web app already uses",
        "  unless a row is marked Newly Developed on New-API only.",
        "",
        "AUTHENTICATION",
        "  Clinic login: POST /api/Account/Login on Old API. JWT in Authorization: Bearer <token>.",
        "  Patient/Doctor mobile OTP: New API /api/Account/RequestOtp + /api/Account/LoginWithOtp (see sheet).",
        "  Old-API JWT and New-API JWT use different signing keys. A token from one host is not valid on the other",
        "  unless that host explicitly accepts it. Prefer logging into the same host you will call.",
        "  Logout: web calls BOTH hosts (Old /api/Account/Logout and New /api/Account/Logout).",
        "",
        "ROLES (high level — row-level Authorization/Role column is source of truth)",
        "  Public / anonymous — no JWT.",
        "  Patient JWT — family, caregiver, consent, patient profile, appointments (as documented).",
        "  Doctor JWT — doctor profile, availability. Doctor MOBILE must NOT implement case-taking APIs.",
        "  Reception — clinic web. Not money roles.",
        "  Account / PharmacyPartner — money dashboards; forbidden from patient PII (403 ForbidMoneyRoles).",
        "  AdminPortal policy — clinical masters mutate + admin ACL.",
        "",
        "CASE-TAKING",
        "  Audio case-taking, clipboard, repertorization, clinical questions, prescription mutate = clinic WEB.",
        "  Do not build these screens on the Doctor mobile app.",
        "",
        "COLUMNS",
        "  Filter API Project = New API vs Old API.",
        "  Filter Mobile Usage for Patient app / Doctor app / Web-only.",
        "  Existing/New/Updated is from week API docs when present; otherwise inferred from which project defines the route.",
        "  Request/response field lists: taken from DTOs or Swagger. If a cell says Not confirmed, do not guess — hit the live API or ask backend.",
        "",
        "THIS SHEET IS COMPLETE FOR EXTERNALLY CALLABLE CONTROLLER ENDPOINTS.",
        "  Template dummy URLs in url_helper.js (fake dashboard widgets) are excluded.",
        "  Internal C# methods without [HttpGet]/[HttpPost]/… are excluded.",
        "",
        "Regenerate: python Documents/_build_mobile_api_reference.py",
        "Also copied under NIGA_NewAPI/ScriptsAndFiles/_build_mobile_api_reference.py",
    ]
    for i, line in enumerate(lines, 2):
        ws[f"A{i}"] = line
        ws[f"A{i}"].alignment = Alignment(wrap_text=True)
        if line and line == line.upper() and not line.startswith(" "):
            ws[f"A{i}"].font = Font(bold=True, color="1F4E79")
    ws.column_dimensions["A"].width = 140
    ws.row_dimensions[1].height = 24


def write_reference(wb, rows: list[list]):
    name = "Mobile_API_Reference"
    if name in wb.sheetnames:
        del wb[name]
    # place after HowTo
    idx = 2 if "Mobile_HowTo" in wb.sheetnames else 1
    ws = wb.create_sheet(name, idx)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    fills = {
        "Newly Developed": PatternFill("solid", fgColor="F4B183"),
        "Updated": PatternFill("solid", fgColor="DDEBF7"),
        "Existing": PatternFill("solid", fgColor="E2EFDA"),
        "Unchanged": PatternFill("solid", fgColor="E2EFDA"),
        "Deprecated": PatternFill("solid", fgColor="F4CCCC"),
        "Replaced": PatternFill("solid", fgColor="FCE4D6"),
        "Blocked": PatternFill("solid", fgColor="FFF2CC"),
    }
    new_fill = PatternFill("solid", fgColor="BDD7EE")
    old_fill = PatternFill("solid", fgColor="FFF2CC")

    ws.freeze_panes = "A2"
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(rows, 2):
        for col, v in enumerate(row, 1):
            cell = ws.cell(i, col, v)
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="Calibri", size=9)
        st = row[4]
        if st in fills:
            ws.cell(i, 5).fill = fills[st]
            ws.cell(i, 11).fill = fills[st]
        if str(row[5]).startswith("New"):
            ws.cell(i, 6).fill = new_fill
        else:
            ws.cell(i, 6).fill = old_fill
    widths = [
        12, 28, 40, 28, 18, 28, 28, 48, 12, 22, 18, 42, 42, 14, 40, 40, 24, 36, 32, 48, 36,
        28, 40, 40, 36, 40, 12, 14, 12, 24, 50, 40, 36, 50, 40, 28, 32, 50, 40, 22,
    ]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(len(rows)+1, 2)}"
    if rows:
        # table names must be unique in workbook
        existing = {t.name for s in wb.worksheets for t in s.tables.values()}
        tname = "MobileApiReference"
        n = 2
        while tname in existing:
            tname = f"MobileApiReference{n}"
            n += 1
        tab = Table(displayName=tname, ref=f"A1:{get_column_letter(len(HEADERS))}{len(rows)+1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)


def patch_readme(wb, n_rows: int):
    if "README" not in wb.sheetnames:
        return
    ws = wb["README"]
    # append note if not already present
    found = False
    for row in ws.iter_rows(min_col=1, max_col=1, max_row=80, values_only=False):
        val = str(row[0].value or "")
        if "Mobile_API_Reference" in val:
            found = True
            break
    last = ws.max_row + 1
    if not found:
        ws[f"A{last}"] = ""
        ws[f"A{last+1}"] = f"Mobile_HowTo + Mobile_API_Reference ({n_rows} endpoints) — complete callable HTTP inventory for the mobile developer. Regenerated {TODAY}."
        ws[f"A{last+1}"].font = Font(bold=True, color="C65911")


def main() -> None:
    import sys

    sys.path.insert(0, str(DOCS))

    new_ctrl = parse_controllers(NEW_CTRL, "New API (nigahomeoAPI)")
    old_ctrl = parse_controllers(OLD_CTRL, "Old API (classic / api)")
    new_sw = parse_swagger(SWAGGER_NEW, "New API (nigahomeoAPI)")
    old_sw = parse_swagger(SWAGGER_OLD, "Old API (classic / api)")
    new_rows = merge_rows(new_ctrl, new_sw, "New API (nigahomeoAPI)")
    old_rows = merge_rows(old_ctrl, old_sw, "Old API (classic / api)")
    docs = parse_api_docs()
    consts = parse_url_helper()
    react = parse_react_usage(consts)
    dto_fields = load_dto_fields()
    sheet_rows = build_sheet_rows(new_rows, old_rows, docs, react, dto_fields)

    if not XLSX.exists():
        raise SystemExit(f"Workbook not found: {XLSX}")
    wb = load_workbook(XLSX)
    write_howto(wb, len(new_rows), len(old_rows), len(sheet_rows))
    write_reference(wb, sheet_rows)
    patch_readme(wb, len(sheet_rows))
    wb.save(XLSX)
    print(f"Wrote {XLSX}")
    print(f"New-API endpoints: {len(new_rows)} (controllers {len(new_ctrl)}, swagger {len(new_sw)})")
    print(f"Old-API endpoints: {len(old_rows)} (controllers {len(old_ctrl)}, swagger {len(old_sw)})")
    print(f"Mobile_API_Reference rows: {len(sheet_rows)}")


if __name__ == "__main__":
    main()
