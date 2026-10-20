# -*- coding: utf-8 -*-
"""Add the mobile clinic routes to the API workbook and point New API samples at the deployed host."""
from __future__ import annotations

from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PATH = r"C:\TufanPowar\OctazenTechnologies\NIGA_NewAPI\ScriptsAndFiles\Homeocentrum_All_New_And_Updated+APIs.xlsx"
BASE = "https://devapi.homeocentrum.com"
API = BASE + "/api"

REPLACEMENTS = (
    ("https://devapi2.homeocentrum.com", BASE),
    ("http://devapi2.homeocentrum.com", BASE),
    ("https://devmaigrationapi.homeocentrum.com", BASE),
    ("http://devmaigrationapi.homeocentrum.com", BASE),
    ("http://127.0.0.1:5002", BASE),
    ("https://127.0.0.1:5002", BASE),
    ("http://localhost:5002", BASE),
)

# method, path, audience (doctor|patient|both), screen, purpose, query, body, response
ROUTES = [
    ("GET", "/api/Public/AppConfig", "both", "First launch / support",
     "Public app settings: site URL, support contacts, and the Razorpay public key id.",
     "No query string.", "No body.",
     '{"success":true,"data":{"razorpayKeyId":"rzp_test_example","siteUrl":"https://devapi.homeocentrum.com","supportPhone":null,"supportWhatsApp":null,"supportEmail":null}}'),
    ("GET", "/api/App/Config", "both", "First launch / support",
     "Same body as GET /api/Public/AppConfig.",
     "No query string.", "No body.",
     '{"success":true,"data":{"razorpayKeyId":"rzp_test_example","siteUrl":"https://devapi.homeocentrum.com","supportPhone":null,"supportWhatsApp":null,"supportEmail":null}}'),
    ("GET", "/api/MobileDoctor/Registration/CheckAvailability", "doctor", "Register clinic",
     "Check whether a username or email is free before doctor registration.",
     "userName, emailId. Either may be omitted. A null side means that parameter was omitted.",
     "No body.",
     '{"success":true,"data":{"userNameAvailable":true,"emailAvailable":true}}'),
    ("GET", "/api/Doctor/PracticeStats", "doctor", "Clinical profile",
     "Practice counts from stored visits, reviews, and signed notes. yearsExperience and punctuality are null.",
     "No query string.", "No body.",
     '{"success":true,"data":{"doctorId":104,"averageRating":null,"reviewCount":0,"completedVisits":28,"cancelledVisits":9,"completionRate":0.757,"lifetimeConsultCount":28,"signedNoteCount":12,"accountAgeDays":24,"yearsExperience":null,"punctuality":null,"punctualityPct":null}}'),
    ("GET", "/api/Doctor/Insights", "doctor", "Insights",
     "Same body as GET /api/Doctor/PracticeStats.",
     "No query string.", "No body.",
     '{"success":true,"data":{"doctorId":104,"completedVisits":28,"yearsExperience":null,"punctuality":null}}'),
    ("GET", "/api/Doctor/FollowUps", "doctor", "Follow-ups",
     "Open or completed follow-up tasks for the signed-in doctor.",
     "status (OPEN), from, to. Dates are yyyy-MM-dd.", "No body.",
     '{"success":true,"data":[{"followUpTaskId":2,"title":"Review cough after 7 days","dueDate":"2026-10-21T00:00:00","status":"OPEN","patientId":5012,"patientAppId":88120,"patientName":"Asha Kulkarni"}]}'),
    ("GET", "/api/Doctor/Activity", "doctor", "Recent activity",
     "Latest 30 appointments for the signed-in doctor.",
     "No query string.", "No body.",
     '{"success":true,"data":[{"kind":"Appointment","at":"2026-10-15T11:00:00","patientAppId":88120,"patientName":"Asha Kulkarni","status":"WAITING","consultMode":"Tele","paymentStatus":"UNPAID"}]}'),
    ("GET", "/api/PatientAppointment/GetAppointmentsByDateRange", "doctor", "Schedule week",
     "Visits for the token user between two dates. Maximum 62 days. consultFee is the doctor's stored tele or in-clinic fee.",
     "from, to. Example from=2026-10-01&to=2026-10-14.", "No body.",
     '{"success":true,"data":[{"patientAppId":88120,"patientName":"Asha Kulkarni","consultFee":500,"currency":null}]}'),
    ("GET", "/api/Doctor/TimeOff", "doctor", "Working hours",
     "Doctor time off. The same list is GET /api/Availability/TimeOff.",
     "No query string.", "No body.",
     '{"success":true,"data":[{"doctorTimeOffId":1,"startsAt":"2026-12-02T09:00:00","endsAt":"2026-12-02T13:00:00","reason":"Camp"}]}'),
    ("POST", "/api/Doctor/TimeOff", "doctor", "Working hours",
     "Save a time-off block. Same body on POST /api/Availability/TimeOff.",
     "No query string.",
     '{"startsAt":"2026-12-01T09:00:00","endsAt":"2026-12-01T13:00:00","reason":"Camp"}',
     '{"success":true,"data":{"doctorTimeOffId":1}}'),
    ("DELETE", "/api/Doctor/TimeOff/{id}", "doctor", "Working hours",
     "Delete one time-off row owned by the doctor. Same on DELETE /api/Availability/TimeOff/{id}.",
     "No query string.", "No body.",
     '{"success":true,"data":{"deleted":true}}'),
    ("GET", "/api/Doctor/BookingRules", "doctor", "Working hours",
     "Gap between slots and how far ahead patients may book. Null until the doctor saves them.",
     "No query string.", "No body.",
     '{"success":true,"data":{"gapMinutes":10,"bookingHorizonDays":14}}'),
    ("PUT", "/api/Doctor/BookingRules", "doctor", "Working hours",
     "Save gap and booking horizon. gapMinutes is 0 to 240. bookingHorizonDays is 1 to 90.",
     "gapMinutes=10&bookingHorizonDays=14", "No body. Send the values as query parameters.",
     '{"success":true,"data":{"gapMinutes":10,"bookingHorizonDays":14}}'),
    ("GET", "/api/Doctor/AppointmentRequests", "doctor", "Requests",
     "Appointment requests sent by patients.",
     "status=PENDING", "No body.",
     '{"success":true,"data":[{"appointmentRequestId":1,"patientId":5012,"patientName":"Asha Kulkarni","requestedDate":"2026-10-21T00:00:00","requestedTime":"10:30","consultMode":"Tele","status":"PENDING","reason":"Fever for two days"}]}'),
    ("POST", "/api/Doctor/AppointmentRequests/{id}/Accept", "doctor", "Requests",
     "Accept a request and create a visit with status NOT ARRIVED, payment UNPAID, channel Request.",
     "No query string.", "No body.",
     '{"success":true,"data":{"status":"ACCEPTED","patientAppId":88121}}'),
    ("POST", "/api/Doctor/AppointmentRequests/{id}/Decline", "doctor", "Requests",
     "Decline an open request.",
     "No query string.", '{"reason":"I am in a camp that day."}',
     '{"success":true,"data":{"status":"DECLINED"}}'),
    ("POST", "/api/Doctor/AppointmentRequests/{id}/Propose", "doctor", "Requests",
     "Propose another date and time.",
     "No query string.", '{"proposedDate":"2026-10-21","proposedTime":"11:00"}',
     '{"success":true,"data":{"status":"PROPOSED"}}'),
    ("GET", "/api/Doctor/CareLinks", "doctor", "Care link",
     "Care-link requests where this doctor is sender or receiver. Same list as GET /api/CareLink. feeInr equals feeAmount. expiry is always null.",
     "No query string.", "No body.",
     '{"success":true,"data":[{"careLinkRequestId":1,"fromDoctorId":104,"toDoctorId":88,"patientId":5012,"patientName":"Asha Kulkarni","note":"Please see this child.","feeAmount":null,"feeInr":null,"opinion":null,"expiry":null,"status":"OPEN"}]}'),
    ("POST", "/api/Doctor/CareLinks", "doctor", "Care link",
     "Create a care-link request. Same body on POST /api/CareLink. feeAmount is only what the sender sends.",
     "No query string.",
     '{"toDoctorId":88,"patientId":5012,"note":"Please see this child.","feeAmount":null}',
     '{"success":true,"data":{"status":"OPEN"}}'),
    ("POST", "/api/Doctor/CareLinks/{id}/Respond", "doctor", "Care link",
     "Receiving doctor accepts or declines while status is OPEN. Same path under /api/CareLink/{id}/Respond.",
     "No query string.", '{"status":"ACCEPTED"}',
     '{"success":true,"data":{"status":"ACCEPTED"}}'),
    ("POST", "/api/Doctor/CareLinks/{id}/Opinion", "doctor", "Care link",
     "Sender or receiver saves an opinion. Same path under /api/CareLink/{id}/Opinion.",
     "No query string.", '{"opinion":"I agree with the current remedy."}',
     '{"success":true,"data":{"careLinkRequestId":1,"opinion":"I agree with the current remedy."}}'),
    ("GET", "/api/Doctor/PharmacyQueries", "doctor", "Pharmacy question",
     "Pharmacy questions for this doctor. Same list as GET /api/PharmacyQuery.",
     "No query string.", "No body.",
     '{"success":true,"data":[{"pharmacyQueryId":1,"patientId":5012,"question":"Is the 30C pack in stock?","reply":null,"status":"OPEN"}]}'),
    ("POST", "/api/Doctor/PharmacyQueries", "doctor", "Pharmacy question",
     "Ask a pharmacy question. Same body on POST /api/PharmacyQuery.",
     "No query string.", '{"patientId":5012,"question":"Is the 30C pack in stock?"}',
     '{"success":true,"data":{"status":"OPEN"}}'),
    ("POST", "/api/Doctor/PharmacyQueries/{id}/Reply", "doctor", "Pharmacy question",
     "Reply to a pharmacy question. Same path under /api/PharmacyQuery/{id}/Reply.",
     "No query string.", '{"reply":"Yes, two bottles."}',
     '{"success":true,"data":{"status":"REPLIED"}}'),
    ("PUT", "/api/Erx/Draft", "doctor", "Prescription",
     "Save remedy lines and notes before sign. remedyId must exist. potencyCode is optional here. Sign still requires a potency on every line.",
     "No query string.",
     '{"patientAppId":88120,"notes":"Review in one week","items":[{"remedyId":4401,"potencyCode":"30C","dose":"4 globules","frequency":"BD","duration":"3 days","instructions":"Away from food"}]}',
     '{"success":true,"data":{"patientAppId":88120,"status":"DRAFT","itemCount":1}}'),
    ("GET", "/api/Notifications/Preferences", "both", "Notification settings",
     "Notification switches for the signed-in user. Defaults before a save: push, appointments, prescriptions, and medicines true.",
     "No query string.", "No body.",
     '{"success":true,"data":{"pushEnabled":true,"smsEnabled":false,"emailEnabled":false,"quietStart":null,"quietEnd":null,"appointments":true,"prescriptions":true,"medicines":true,"marketing":false,"quietHoursEnabled":false,"reminderOffsets":null}}'),
    ("PUT", "/api/Notifications/Preferences", "both", "Notification settings",
     "Save notification switches.",
     "No query string.",
     '{"pushEnabled":true,"smsEnabled":false,"emailEnabled":false,"quietStart":"22:00:00","quietEnd":"07:00:00","appointments":true,"prescriptions":true,"medicines":true,"marketing":false,"quietHoursEnabled":true,"reminderOffsets":"60,1440"}',
     '{"success":true,"data":{"pushEnabled":true,"quietHoursEnabled":true,"reminderOffsets":"60,1440"}}'),
    ("POST", "/api/Notifications/MarkAllRead", "both", "Notifications",
     "Mark every notification for this user as read.",
     "No query string.", "No body.",
     '{"success":true,"data":{"marked":4}}'),
    ("POST", "/api/PatientAppointment/SwitchToTele", "both", "Appointment",
     "Patient or owning doctor switches an in-clinic visit to tele.",
     "patientAppId=88120. A JSON body with patientAppId is also accepted.", "No body when the query parameter is sent.",
     '{"success":true,"data":{"patientAppId":88120,"consultMode":"Tele","isTele":true}}'),
    ("POST", "/api/Tele/Sessions/{id}/RunningLate", "both", "Video call",
     "Patient or owning doctor records that they are running late.",
     "minutes=10. A JSON body with minutes is also accepted.", "No body when the query parameter is sent.",
     '{"success":true,"data":{"teleSessionId":15,"code":"RUNNING_LATE","detail":"10 minutes"}}'),
    ("POST", "/api/Patient/Visits/{id}/CheckIn", "patient", "Clinic check-in",
     "Today's in-clinic visit only. Status is EARLY, WAITING, or READY. Tele visits are rejected.",
     "status=WAITING. A JSON body with status is also accepted.", "No body when the query parameter is sent.",
     '{"success":true,"data":{"patientAppId":88120,"status":"WAITING","queuePosition":2}}'),
    ("GET", "/api/Patient/Reminders", "patient", "Reminders",
     "Reminders for the signed-in patient. kind, timeOfDay, weekdays, and memberPatientId are null when they were not saved.",
     "No query string.", "No body.",
     '{"success":true,"data":[{"patientReminderId":3,"patientId":5012,"title":"Evening dose","dueAt":"2026-10-16T20:00:00","isDone":false,"kind":"MEDICINE","timeOfDay":"20:00","weekdays":"Mon,Wed","memberPatientId":null}]}'),
    ("POST", "/api/Patient/Reminders", "patient", "Reminders",
     "Create a reminder.",
     "No query string.",
     '{"title":"Evening dose","dueAt":"2026-10-16T20:00:00","kind":"MEDICINE","timeOfDay":"20:00","weekdays":"Mon,Wed","memberPatientId":null}',
     '{"success":true,"data":{"patientReminderId":3}}'),
    ("PUT", "/api/Patient/Reminders/{id}", "patient", "Reminders",
     "Update a reminder owned by the signed-in user.",
     "No query string.",
     '{"title":"Evening dose","dueAt":"2026-10-16T20:00:00","isDone":true,"kind":"MEDICINE","timeOfDay":"20:00","weekdays":"Mon,Wed","memberPatientId":null}',
     '{"success":true,"data":{"patientReminderId":3}}'),
    ("DELETE", "/api/Patient/Reminders/{id}", "patient", "Reminders",
     "Delete a reminder owned by the signed-in user.",
     "No query string.", "No body.",
     '{"success":true,"data":{"deleted":true}}'),
    ("POST", "/api/Patient/AppointmentRequests", "patient", "Book a visit",
     "Ask the doctor for a visit. This does not create the appointment until the doctor accepts.",
     "No query string.",
     '{"doctorId":104,"requestedDate":"2026-10-20","requestedTime":"10:30","consultMode":"Tele","reason":"Fever for two days"}',
     '{"success":true,"data":{"status":"PENDING"}}'),
]

NOTES = {
    "/api/earnings/summary": " unpaidReceivableAmount is the stored consult fee times unpaid, non-cancelled visits. It is not pendingPayoutAmount. currency is null when neither a payment nor the doctor fee currency is stored.",
    "/api/doctormobile/context/": " Response includes consultFee, currency, platformFee, totalAmount, and paidAt. platformFee is null unless the captured payment is greater than the consult fee.",
    "/api/profile/me": " registrationNo repeats passingCertNo. registrationValidUntil, issuingCouncil, and yearsExperience are null. photoUrl is absolute when a photo is stored.",
    "/api/profile/me/credentials": " Each document includes validUntil, verified, and issuingCouncil. Those three are null. Download with downloadUrl.",
    "/api/patient/visits": " List items include consultFee, currency, clinicPhone, clinicAddress, and reasonText. reasonText is null on the list and set on the visit detail from the chief complaint.",
    "/api/reviews/mine": " A patient token returns that patient's reviews with doctorName and createdAt. A doctor token returns the doctor's reviews.",
    "/api/pharmacy/sellers": " priceTotal, deliveryFee, etaMinutes, fulfilPercent, distanceKm, and serviceable are null. Those facts are not stored.",
    "getallcases": " search filters patient name, mobile, and email.",
    "getcountapp": " Response adds unpaidAmount, patientAppCancelled, unpaidAppointmentCount, and currency from stored visits and the doctor fee.",
    "/api/availability/me": " Response includes gapMinutes and bookingHorizonDays. Both are null until saved.",
}


def norm(method: str, endpoint: str) -> tuple[str, str]:
    path = endpoint.split("?")[0].rstrip("/").lower()
    path = path.replace("{id:int}", "{id}").replace("{patientappid:int}", "{id}")
    return method.upper(), path


def replace_hosts(wb) -> int:
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                text = cell.value
                new = text
                for old, repl in REPLACEMENTS:
                    new = new.replace(old, repl)
                if new != text:
                    cell.value = new
                    n += 1
    return n


def append_note(wb) -> int:
    n = 0
    for name in ("All_APIs", "New_APIs", "Doctor_Mobile", "Patient_App", "S4_Week4", "S5_Week5", "Week4_Week5_Developer"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if "API Endpoint (as in MD)" not in headers or "What was developed this week" not in headers:
            continue
        ep_i = headers.index("API Endpoint (as in MD)")
        note_i = headers.index("What was developed this week")
        for row in ws.iter_rows(min_row=2):
            ep = str(row[ep_i].value or "").lower()
            for key, extra in NOTES.items():
                if key in ep and extra.strip() not in str(row[note_i].value or ""):
                    row[note_i].value = (str(row[note_i].value or "").rstrip() + extra)
                    n += 1
                    break
    return n


def existing_keys(ws) -> set[tuple[str, str]]:
    keys = set()
    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if row[1] and row[2]:
            endpoint = str(row[1])
            if " " in endpoint:
                method, path = endpoint.split(" ", 1)
            else:
                method, path = str(row[2]), endpoint
            keys.add(norm(method, path))
    return keys


def catalog_row(sr: int, method: str, path: str, audience: str, screen: str, purpose: str, query: str, body: str, response: str) -> list:
    patient = "Yes" if audience in ("patient", "both") else "No"
    doctor = "Yes" if audience in ("doctor", "both") else "No"
    token = "no" if path in ("/api/Public/AppConfig", "/api/App/Config", "/api/MobileDoctor/Registration/CheckAvailability") else "yes"
    who = {"doctor": "Doctor mobile app", "patient": "Patient app", "both": "Patient app and Doctor mobile app"}[audience]
    user = {"doctor": "Doctor", "patient": "Patient", "both": "Patient, Doctor"}[audience]
    full = BASE + path
    sample_req = f"{method} {full}\nAuthorization: Bearer <data.token>\nContent-Type: application/json\n\nQuery: {query}\n\nBody:\n{body}"
    if token == "no":
        sample_req = f"{method} {full}\nNo Authorization header.\n\nQuery: {query}\n\nBody:\n{body}"
    login = (
        "POST " + API + "/Account/Login\n"
        "Content-Type: application/json\n"
        '{"userName":"<username from Users & Login Details>","password":"<password from that sheet>"}\n'
        "Use data.token as Authorization: Bearer <data.token>."
    )
    row = [None] * 72
    values = {
        1: sr,
        2: f"{method} {path}",
        3: method,
        4: "New API developed",
        5: "Mobile clinic route added 15-Oct-2026. " + purpose,
        6: "Mobile clinic",
        7: user,
        8: "S5_Week5",
        9: who,
        10: patient,
        11: doctor,
        12: "No",
        13: "Yes" if audience == "both" else "No",
        14: "New-API",
        15: API,
        16: token,
        17: "MOBILE-2026-10-15",
        18: purpose,
        19: sample_req,
        20: "Fictitious sample. Names and ids in the sample are not live patients.",
        21: response,
        22: "Call this on " + API + ". Hide any field that is null. Do not invent a fee, phone, year, or rating.",
        23: "No token" if token == "no" else "Bearer JWT",
        24: "400 validation, 401 no/expired token, 403 wrong role, 404 missing row, 409 already signed. Body is success, code, message.",
        25: screen,
        26: "Do not call the Old API for this path.",
        27: f"MOB-{sr:03d}",
        28: "NIGA_NewAPI mobile controllers, 15-Oct-2026",
        29: BASE,
        30: full,
        31: "application/json when a body is sent",
        32: "Accept: application/json" if token == "no" else "Accept: application/json\nAuthorization: Bearer <data.token>",
        33: "No login." if token == "no" else login,
        34: "Path id is the number in the URL." if "{id}" in path else "No path id. The URL path is fixed.",
        35: query,
        36: body,
        37: body,
        38: response,
        39: '{"success":false,"code":"VALIDATION","message":"The message names the missing field."}',
        40: who + ". Screen: " + screen,
        41: "Sample people in this row are fictitious. Base URL is " + BASE + ".",
        42: path.strip("/").split("/")[-1],
        43: "MobileDoctorController.cs / MobilePatientController.cs / MobileSharedController.cs",
        44: full,
        45: "New API",
        46: "Yes",
        47: "No" if token == "no" else "Yes",
        48: token,
        49: "Active",
        50: "Not Applicable",
        51: "See Request body fields.",
        52: "See Request body fields.",
        53: "None",
        54: query + " " + body,
        55: "See Sample response.",
        56: "200, 400, 401, 403, 404",
        57: "Show the API message. 401: clear the token and open login. 403: do not retry.",
        58: "No token." if token == "no" else "Authorization: Bearer <data.token> from POST /api/Account/Login on " + BASE + ".",
        59: user,
        60: "See the mobile integration guide for the screen order.",
        61: "No",
        62: query,
        63: "Send dates as yyyy-MM-dd or yyyy-MM-ddTHH:mm:ss. No timezone conversion is applied.",
        64: "Not Applicable",
        65: "Not Applicable",
        66: "Not Applicable",
        67: purpose,
        68: "Use this path on " + API + ".",
        69: "Added 15-Oct-2026.",
        70: "One call per screen load. Do not poll faster than the screen already does.",
        71: "Use HTTPS. Do not log the token.",
        72: screen,
    }
    for col, value in values.items():
        row[col - 1] = value
    return row


def contract_row(method: str, path: str, audience: str, screen: str, purpose: str, query: str, body: str, response: str) -> list:
    patient = "Yes" if audience in ("patient", "both") else "No"
    doctor = "Yes" if audience in ("doctor", "both") else "No"
    token = path not in ("/api/Public/AppConfig", "/api/App/Config", "/api/MobileDoctor/Registration/CheckAvailability")
    return [
        path.strip("/").split("/")[-1],
        purpose,
        "Mobile clinic",
        "MobileDoctorController.cs / MobilePatientController.cs / MobileSharedController.cs",
        method,
        path,
        API,
        BASE + path,
        "New API",
        "New API developed",
        "Yes",
        patient,
        doctor,
        "No",
        screen + ". " + purpose,
        {"doctor": "Doctor", "patient": "Patient", "both": "Patient, Doctor"}[audience],
        "No" if not token else "Yes",
        "No token" if not token else "Bearer JWT from POST /api/Account/Login",
        "No",
        "Not Applicable",
        body,
        query,
        "None",
        "{id}" if "{id}" in path else "None",
        query,
        "200, 400, 401, 403, 404",
        "No",
        "No",
        purpose,
        "Do not call the Old API for this path.",
        screen,
        response,
    ]


def extend_filter(ws) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.freeze_panes = "A2"


def main() -> None:
    wb = load_workbook(PATH)
    replaced = replace_hosts(wb)
    noted = append_note(wb)

    all_ws = wb["All_APIs"]
    keys = existing_keys(all_ws)
    added = []
    for item in ROUTES:
        method, path = item[0], item[1]
        if norm(method, path) in keys:
            continue
        added.append(item)
        keys.add(norm(method, path))

    targets = {
        "All_APIs": lambda a: True,
        "New_APIs": lambda a: True,
        "Doctor_Mobile": lambda a: a in ("doctor", "both"),
        "Patient_App": lambda a: a in ("patient", "both"),
        "S5_Week5": lambda a: True,
        "Week4_Week5_Developer": lambda a: True,
    }
    for item in added:
        audience = item[2]
        for sheet, pred in targets.items():
            if sheet not in wb.sheetnames or not pred(audience):
                continue
            ws = wb[sheet]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            if not headers or headers[1] != "API Endpoint (as in MD)":
                continue
            sr = ws.max_row
            row = catalog_row(sr, *item)
            ws.append(row[: len(headers)])
            for cell in ws[ws.max_row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    contract = wb["Mobile_API_Contract"]
    contract_keys = set()
    for row in contract.iter_rows(min_row=2, max_col=6, values_only=True):
        if row[4] and row[5]:
            contract_keys.add(norm(str(row[4]), str(row[5])))
    for item in added:
        if norm(item[0], item[1]) in contract_keys:
            continue
        contract.append(contract_row(*item))

    for name in list(targets) + ["Mobile_API_Contract"]:
        if name in wb.sheetnames:
            extend_filter(wb[name])

    readme = wb["README"]
    today = date.today().strftime("%d-%b-%Y")
    readme.append([f"Updated {today}. Mobile base URL for every New API call: {API}"])
    readme.append(["Old API host stays https://devapi1.homeocentrum.com/api. Mobile apps call the New API host only."])
    readme.append(["Mobile clinic routes added on this date are marked Mobile clinic in MD section. Sample JSON on those rows is fictitious."])
    readme.append(["Login: POST " + API + "/Account/Login with the username and password on Users & Login Details. Send data.token as Authorization: Bearer."])
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in (readme.cell(readme.max_row - 3, 1),):
        cell.fill = fill
        cell.font = font

    wb.save(PATH)
    print(f"replaced_cells={replaced}")
    print(f"notes_appended={noted}")
    print(f"routes_added={len(added)}")
    for item in added:
        print(item[0], item[1])


if __name__ == "__main__":
    main()
