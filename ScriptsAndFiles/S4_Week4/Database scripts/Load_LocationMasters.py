"""
Reload Country, State, District, City, and PinCode masters from the three
location Excel files.

Country Code.xlsx        -> CountryMaster (calling code stored with a leading +)
Country VS State.xlsx    -> StateMaster under the matching country
India City Pincode.xlsx  -> District, City (Location), PinCode for India
"""
import datetime
import unicodedata
from collections import OrderedDict

import openpyxl
import pyodbc

DATA = r"C:\TufanPowar\OctazenTechnologies\Documents\LocationMasters\Data"
COUNTRY_FILE = DATA + r"\Country Code.xlsx"
STATE_FILE = DATA + r"\Country VS State.xlsx"
PIN_FILE = DATA + r"\India City Pincode.xlsx"

ENTERED_BY = "Tufan Powar"
ENTERED_DATE = datetime.datetime(2026, 10, 2)

CONN = (
    "Driver={ODBC Driver 17 for SQL Server};"
    "Server=localhost\\MSSQLSERVER25;"
    "Database=HomeoCentrum_Dev;"
    "Trusted_Connection=yes;"
)

# Pincode workbook uses older names. State workbook has the current India name.
STATE_ALIASES = {
    "orissa": "odisha",
    "dadra and nagar haveli": "dadra and nagar haveli and daman and diu",
    "daman and diu": "dadra and nagar haveli and daman and diu",
}


def norm(value):
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.lower().replace("&", "and").split())


def calling_code(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.startswith("+"):
        return text
    return "+" + text


def pin_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text or None


def load_countries():
    wb = openpyxl.load_workbook(COUNTRY_FILE, read_only=True, data_only=True)
    ws = wb["Country Codes"]
    rows = []
    seen = set()
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 0 or not row or not row[0]:
            continue
        name = str(row[0]).strip()
        key = norm(name)
        if key in seen:
            continue
        seen.add(key)
        iso2 = str(row[1]).strip().upper() if row[1] else None
        iso3 = str(row[2]).strip().upper() if row[2] else None
        rows.append((name, calling_code(row[3]), iso2, iso3, key))
    wb.close()
    return rows


def load_states(country_keys):
    wb = openpyxl.load_workbook(STATE_FILE, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    headers = None
    buckets = OrderedDict()
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 0:
            headers = list(row)
            continue
        for column, header in enumerate(headers):
            if not header:
                continue
            country_key = norm(header)
            if country_key not in country_keys:
                continue
            value = row[column] if row and column < len(row) else None
            if not value:
                continue
            name = str(value).strip()
            if not name:
                continue
            states = buckets.setdefault(country_key, OrderedDict())
            states.setdefault(norm(name), name)
    wb.close()
    return buckets


def load_india(india_by_norm):
    wb = openpyxl.load_workbook(PIN_FILE, read_only=True, data_only=True)
    ws = wb.active
    districts = OrderedDict()
    cities = OrderedDict()
    pins = set()
    skipped = OrderedDict()
    used = 0
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 0 or not row or len(row) < 4:
            continue
        raw_state, raw_district, raw_city, raw_pin = row[:4]
        if not raw_state or not raw_district or not raw_city:
            continue
        state_key = norm(raw_state)
        if state_key == "state":
            continue
        state_key = STATE_ALIASES.get(state_key, state_key)
        state_name = india_by_norm.get(state_key)
        if not state_name:
            skipped[str(raw_state).strip()] = skipped.get(str(raw_state).strip(), 0) + 1
            continue
        district_name = str(raw_district).strip()
        city_name = str(raw_city).strip()
        pin = pin_text(raw_pin)
        if not district_name or not city_name or not pin:
            continue
        if norm(district_name) == "district" and norm(city_name) == "location":
            continue
        district_key = (norm(state_name), norm(district_name))
        districts.setdefault(district_key, (state_name, district_name))
        city_key = (norm(state_name), norm(district_name), norm(city_name))
        cities.setdefault(city_key, (state_name, district_name, city_name))
        pins.add((city_key, pin))
        used += 1
    wb.close()
    return districts, cities, pins, used, skipped


def executemany(cursor, sql, rows, size=2000):
    cursor.fast_executemany = True
    batch = []
    for row in rows:
        batch.append(row)
        if len(batch) >= size:
            cursor.executemany(sql, batch)
            batch.clear()
    if batch:
        cursor.executemany(sql, batch)


def main():
    countries = load_countries()
    country_keys = {row[4] for row in countries}
    states_by_country = load_states(country_keys)
    india_states = states_by_country.get("india", {})
    india_by_norm = {key: name for key, name in india_states.items()}
    districts, cities, pins, used_rows, skipped = load_india(india_by_norm)

    connection = pyodbc.connect(CONN)
    connection.autocommit = False
    cursor = connection.cursor()
    try:
        cursor.execute(
            """
            IF COL_LENGTH(N'dbo.CountryMaster', N'Iso2Code') IS NULL
                ALTER TABLE dbo.CountryMaster ADD Iso2Code nvarchar(2) NULL;
            IF COL_LENGTH(N'dbo.CountryMaster', N'Iso3Code') IS NULL
                ALTER TABLE dbo.CountryMaster ADD Iso3Code nvarchar(3) NULL;
            """
        )
        cursor.execute("UPDATE dbo.Doctor SET CityId = NULL, DistrictId = NULL, StateId = NULL, CountryId = NULL")
        cursor.execute("UPDATE dbo.Patient SET StateId = NULL, CountryId = NULL")
        cursor.execute("UPDATE dbo.UserMaster SET StateId = NULL, CountryId = NULL")
        cursor.execute("DELETE FROM dbo.PinCodeMaster")
        cursor.execute("DELETE FROM dbo.CityMaster")
        cursor.execute("DELETE FROM dbo.DistrictMaster")
        cursor.execute("DELETE FROM dbo.StateMaster")
        cursor.execute("DELETE FROM dbo.CountryMaster")
        cursor.execute("DBCC CHECKIDENT ('dbo.PinCodeMaster', RESEED, 0)")
        cursor.execute("DBCC CHECKIDENT ('dbo.CityMaster', RESEED, 0)")
        cursor.execute("DBCC CHECKIDENT ('dbo.DistrictMaster', RESEED, 0)")
        cursor.execute("DBCC CHECKIDENT ('dbo.StateMaster', RESEED, 0)")
        cursor.execute("DBCC CHECKIDENT ('dbo.CountryMaster', RESEED, 0)")

        executemany(
            cursor,
            """
            INSERT INTO dbo.CountryMaster
                (CountryName, CountryCode, Iso2Code, Iso3Code, EnteredBy, EnteredDate, DeleteStatus)
            VALUES (?, ?, ?, ?, ?, ?, 0)
            """,
            [
                (name, code, iso2, iso3, ENTERED_BY, ENTERED_DATE)
                for name, code, iso2, iso3, _key in countries
            ],
        )
        cursor.execute("SELECT CountryId, CountryName FROM dbo.CountryMaster")
        country_ids = {norm(name): country_id for country_id, name in cursor.fetchall()}

        state_rows = []
        for country_key, states in states_by_country.items():
            country_id = country_ids[country_key]
            for state_name in states.values():
                state_rows.append((state_name, country_id, ENTERED_BY, ENTERED_DATE))
        executemany(
            cursor,
            """
            INSERT INTO dbo.StateMaster
                (StateName, CountryId, EnteredBy, EnteredDate, DeleteStatus)
            VALUES (?, ?, ?, ?, 0)
            """,
            state_rows,
        )
        cursor.execute("SELECT StateId, StateName, CountryId FROM dbo.StateMaster")
        india_id = country_ids["india"]
        state_ids = {
            norm(name): state_id
            for state_id, name, country_id in cursor.fetchall()
            if country_id == india_id
        }

        executemany(
            cursor,
            """
            INSERT INTO dbo.DistrictMaster
                (DistrictName, StateId, EnteredBy, EnteredDate, DeleteStatus)
            VALUES (?, ?, ?, ?, 0)
            """,
            [
                (district_name, state_ids[norm(state_name)], ENTERED_BY, ENTERED_DATE)
                for state_name, district_name in districts.values()
            ],
        )
        cursor.execute(
            """
            SELECT d.DistrictId, d.DistrictName, s.StateName
            FROM dbo.DistrictMaster d
            JOIN dbo.StateMaster s ON s.StateId = d.StateId
            WHERE s.CountryId = ?
            """,
            india_id,
        )
        district_ids = {
            (norm(state_name), norm(district_name)): district_id
            for district_id, district_name, state_name in cursor.fetchall()
        }

        executemany(
            cursor,
            """
            INSERT INTO dbo.CityMaster
                (CityName, DistrictId, EnteredBy, EnteredDate, DeleteStatus)
            VALUES (?, ?, ?, ?, 0)
            """,
            [
                (city_name, district_ids[(norm(state_name), norm(district_name))], ENTERED_BY, ENTERED_DATE)
                for state_name, district_name, city_name in cities.values()
            ],
        )
        cursor.execute(
            """
            SELECT c.CityId, c.CityName, d.DistrictName, s.StateName
            FROM dbo.CityMaster c
            JOIN dbo.DistrictMaster d ON d.DistrictId = c.DistrictId
            JOIN dbo.StateMaster s ON s.StateId = d.StateId
            WHERE s.CountryId = ?
            """,
            india_id,
        )
        city_ids = {
            (norm(state_name), norm(district_name), norm(city_name)): city_id
            for city_id, city_name, district_name, state_name in cursor.fetchall()
        }

        executemany(
            cursor,
            """
            INSERT INTO dbo.PinCodeMaster
                (PinCode, CityId, EnteredBy, EnteredDate, DeleteStatus)
            VALUES (?, ?, ?, ?, 0)
            """,
            [
                (pin, city_ids[city_key], ENTERED_BY, ENTERED_DATE)
                for city_key, pin in pins
            ],
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()

    print("countries", len(countries))
    print("states", sum(len(items) for items in states_by_country.values()))
    print("districts", len(districts))
    print("cities", len(cities))
    print("pincodes", len(pins))
    print("india pin rows used", used_rows)
    print("unmatched pin states", dict(skipped))


if __name__ == "__main__":
    main()
